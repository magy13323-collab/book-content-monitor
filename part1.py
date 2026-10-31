"""Offline, lossless processing of the first worksheet. No page fetching."""

import argparse
from collections import defaultdict
from dataclasses import dataclass, field
from difflib import SequenceMatcher
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from urllib.parse import unquote, urlsplit, urlunsplit

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = ("Правообладатель", "Автор", "Название", "Тип контента", "url", "Сайт")
PUNCTUATION = str.maketrans({
    "«": '"', "»": '"', "“": '"', "”": '"', "„": '"',
    "‘": "'", "’": "'", "‐": "-", "‑": "-", "–": "-", "—": "-", "−": "-",
})
RUSSIAN = str.maketrans(dict(zip(
    "абвгдежзийклмнопрстуфхцчшщъыьэюя",
    ("a", "b", "v", "g", "d", "e", "zh", "z", "i", "i", "k", "l", "m", "n",
     "o", "p", "r", "s", "t", "u", "f", "kh", "ts", "ch", "sh", "shch", "",
     "y", "", "e", "iu", "ia"),
)))


def normalize_identity(value):
    """Identity key: preserve е/ё, punctuation, numbers and author order."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    text = re.sub(r"\s+", " ", text.translate(PUNCTUATION))
    return text.strip(" ;*")


def normalize_text(value):
    """Search similarity only; е/ё folding does not establish identity."""
    return normalize_identity(value).replace("ё", "е")


def normalize_url(value):
    """Only scheme/host case and default ports; path/query/fragment stay intact."""
    text = "" if value is None else str(value).strip()
    if re.search(r"\s", text):
        return text  # urlsplit silently removes internal tabs/newlines; keep them for review.
    try:
        parts = urlsplit(text)
        if not parts.hostname or parts.scheme.lower() not in {"http", "https"}:
            return text
        scheme = parts.scheme.lower()
        host = parts.hostname.lower()
        if ":" in host:  # IPv6
            host = f"[{host}]"
        port = parts.port
        if port is not None and (scheme, port) not in {("http", 80), ("https", 443)}:
            host += f":{port}"
        if "@" in parts.netloc:
            host = parts.netloc.rsplit("@", 1)[0] + "@" + host
        result = urlunsplit((scheme, host, parts.path, parts.query, parts.fragment))
        # urlunsplit drops empty delimiters; conservatively preserve them too.
        if "?" in text.split("#", 1)[0] and not parts.query:
            head, separator, fragment = result.partition("#")
            result = head + "?" + (separator + fragment if separator else "")
        if text.endswith("#") and not result.endswith("#"):
            result += "#"
        return result
    except ValueError:
        return text


def valid_url(value):
    try:
        parts = urlsplit(value)
        return (parts.scheme in {"http", "https"} and bool(parts.hostname)
                and parts.port != 0 and not re.search(r"\s", value))
    except ValueError:
        return False


def group_id(prefix, key):
    payload = json.dumps(key, ensure_ascii=False, default=str, sort_keys=True)
    return prefix + "_" + hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


@dataclass
class Record:
    source_file: str
    source_sheet: str
    source_row: int
    raw: tuple
    author: str
    title: str
    url: str
    content_type: str
    record_group_id: str = ""
    url_group_id: str = ""
    book_id: str = ""
    issues: set = field(default_factory=set)
    review_ids: list = field(default_factory=list)


@dataclass
class Analysis:
    source_path: Path
    records: list
    duplicates: list = field(default_factory=list)
    books: list = field(default_factory=list)
    review: list = field(default_factory=list)


def buckets(records, key):
    result = defaultdict(list)
    for record in records:
        result[key(record)].append(record)
    return result


def add_review(analysis, code, records, reason, needed_data):
    records = sorted(records, key=lambda record: record.source_row)
    rows = [record.source_row for record in records]
    review_id = group_id("review", (code, records[0].source_sheet, rows))
    analysis.review.append({
        "review_id": review_id, "code": code, "source_rows": rows,
        "reason": reason, "needed_data": needed_data,
        "book_ids": sorted({record.book_id for record in records}),
    })
    for record in records:
        record.issues.add(code)
        record.review_ids.append(review_id)


def author_signature(author):
    # Used for review only, never to merge books or remove aliases.
    core = re.sub(r"\([^)]*\)", " ", author)
    return tuple(sorted(re.findall(r"[^\W_]+", core)))


def slug_tokens(text):
    text = unicodedata.normalize("NFKD", normalize_text(text).translate(RUSSIAN))
    text = "".join(character for character in text if not unicodedata.combining(character))
    # Common alternative transliterations, used exclusively for review hints.
    for original, replacement in (("kh", "h"), ("yu", "u"), ("ju", "u"), ("iu", "u"),
                                  ("ya", "a"), ("ja", "a"), ("ia", "a"), ("ts", "c"),
                                  ("y", "i")):
        text = text.replace(original, replacement)
    return re.findall(r"[a-z]+", text)


def url_title_checks(analysis, record):
    """Advisory URL hints; a slug is not evidence about the actual page."""
    path = unquote(urlsplit(record.url).path)
    slug = path.rstrip("/").rsplit("/", 1)[-1]
    slug = re.sub(r"\.(?:html?|php|aspx?)$", "", slug.casefold())
    url_numbers = {
        label: int(number) for label, number in re.findall(
            r"(?:^|[-_.])(tom|kniga|chast|cast|ch)[-_.]?(\d+)\b", slug
        )
    }
    conflicts = []
    for label, number in re.findall(r"\b(том|книга|часть|ч\.)\s*(\d+)\b", record.title):
        aliases = {"том": ("tom",), "книга": ("kniga",),
                   "часть": ("chast", "cast", "ch"), "ч.": ("chast", "cast", "ch")}[label]
        found = [url_numbers[alias] for alias in aliases if alias in url_numbers]
        if any(value != int(number) for value in found):
            conflicts.append(f"{label} {number} в названии, {found} в URL")
    hash_number = re.search(r"#\s*(\d+)\b", record.title)
    slug_number = re.search(r"(?:^|[-_])(\d{1,3})(?=[-_.]|$)", slug)
    if hash_number and slug_number and int(hash_number[1]) != int(slug_number[1]):
        conflicts.append(f"№{hash_number[1]} в названии, №{slug_number[1]} в URL")
    if conflicts:
        add_review(analysis, "url_number_conflict", [record],
                   "; ".join(conflicts) + ". URL — лишь косвенный признак.",
                   "Официальная карточка: серия, номер книги/тома/части; содержимое страницы.")

    # Opaque IDs and script endpoints do not contain useful title hints.
    tokens = [token for token in slug_tokens(slug) if len(token) >= 3]
    if not tokens or slug in {"index", "viewtopic"} or not re.search(r"[-_]", slug):
        return
    title_tokens = [token for token in slug_tokens(record.title) if len(token) >= 3]
    if not title_tokens:
        return
    matches = [max(SequenceMatcher(None, token, other).ratio() for other in tokens)
               for token in title_tokens]
    coverage = sum(len(token) for token, score in zip(title_tokens, matches) if score >= .78)
    coverage /= sum(map(len, title_tokens))
    compact_similarity = SequenceMatcher(None, "".join(slug_tokens(record.title)),
                                         "".join(slug_tokens(slug))).ratio()
    if coverage < .5 and compact_similarity < .88:
        add_review(analysis, "title_url_hint", [record],
                   f"Слабое сходство названия с текстом URL (покрытие слов {coverage:.0%}). "
                   "Возможны другое произведение, иной подзаголовок или другая транслитерация.",
                   "Официальное название/варианты названия, автор и содержимое страницы.")
    elif len(title_tokens) == 1 and len(tokens) >= 5:
        add_review(analysis, "short_title_url", [record],
                   "Короткое название совпадает только с частью длинного URL; "
                   "этого недостаточно для идентификации книги и автора.",
                   "Полное название, автор, ISBN/ID каталога и содержимое страницы.")


def find_review_cases(analysis):
    records = analysis.records
    for record in records:
        missing = [HEADERS[index] for index, value in enumerate(record.raw)
                   if value is None or not str(value).strip()]
        if missing:
            add_review(analysis, "missing_required", [record],
                       "Не заполнены поля: " + ", ".join(missing), "Недостающие исходные значения.")
        if not record.author or not record.title or not record.content_type:
            record.issues.add("incomplete_book_key")
        for name, value in (("author", record.raw[1]), ("title", record.raw[2])):
            if value is not None and re.match(r"^[\s;*]|[;*]\s*$", str(value)):
                # Only visible punctuation cleanup needs an annotation, not ordinary casefold.
                if str(value).strip().startswith((";", "*")) or str(value).strip().endswith((";", "*")):
                    record.issues.add(name + "_edge_noise")
        if not valid_url(record.url):
            add_review(analysis, "invalid_url", [record],
                       "URL отсутствует или не является корректным абсолютным HTTP(S)-адресом.",
                       "Корректный URL источника.")
            continue
        parts = urlsplit(record.url)
        site = str(record.raw[5] or "").strip().lower()
        if "://" in site:
            try:
                site = urlsplit(site).hostname
            except ValueError:
                site = ""
        if site != parts.hostname:
            add_review(analysis, "site_host_mismatch", [record],
                       f"Поле «Сайт» ({record.raw[5]}) отличается от хоста URL ({parts.hostname}).",
                       "Уточнить: поле содержит точный хост или группу сайтов/зеркал.")
        if re.search(r"/(?:tag|tags|author|authors)/", parts.path, re.I):
            add_review(analysis, "listing_url", [record],
                       "Путь URL похож на страницу автора/тега, а не отдельной книги.",
                       "Прямая ссылка на книгу; наличие книги и её контента на странице.")
        url_title_checks(analysis, record)

    for group in buckets(records, lambda record: (
        normalize_text(record.author), normalize_text(record.title), record.content_type
    )).values():
        if len({(record.author, record.title) for record in group}) > 1:
            add_review(analysis, "yo_variant", group,
                       "Названия или имена авторов различаются е/ё. Кандидаты сохранены отдельно; "
                       "сходство написания не подтверждает и не опровергает тождество произведений.",
                       "ISBN/ID произведения или подтверждённое библиографическое написание.")

    for group in buckets(records, lambda record: record.book_id).values():
        title = group[0].title
        if len(title) < 10 or not title.isalpha():
            continue
        for record in group:
            if not valid_url(record.url):
                continue
            slug = unquote(urlsplit(record.url).path).rstrip("/").rsplit("/", 1)[-1]
            tokens = [token for token in slug_tokens(slug) if token not in {"book", "books", "kniga"}]
            similarity = SequenceMatcher(None, "".join(slug_tokens(title)), "".join(tokens)).ratio()
            if len(tokens) >= 2 and similarity >= .8:
                add_review(analysis, "unsegmented_title", group,
                           "Название записано слитно, а похожий текст URL содержит разделённые слова. "
                           "Возможны пропущенные пробелы; название сохранено без исправлений.",
                           "Официальное написание названия и ID книги; URL не доказывает ошибку.")
                break

    for group in buckets(records, lambda record: record.title).values():
        if group[0].title and len({record.author for record in group}) > 1:
            add_review(analysis, "title_author_conflict", group,
                       "Одно нормализованное название указано у разных авторов. Книги не объединены.",
                       "ISBN/ID и официальные карточки; подтверждённые сведения об авторстве.")

    for group in buckets(records, lambda record: (record.author, record.title)).values():
        if len({record.content_type for record in group}) > 1:
            add_review(analysis, "content_type_conflict", group,
                       "У одного автора и названия разные типы контента. Кандидаты разделены.",
                       "Идентификаторы изданий и правило связи электронной/аудиоверсии.")

    for signature, group in buckets(records, lambda record: author_signature(record.author)).items():
        if signature and len({record.author for record in group}) > 1:
            add_review(analysis, "author_variant", group,
                       "Возможная перестановка имени или вариант с псевдонимом в скобках. "
                       "Имена сохранены; это не подтверждение одного автора.",
                       "ID автора или подтверждённый справочник имён и псевдонимов.")

    for group in buckets(records, lambda record: (
        record.author, record.title.replace(" ", ""), record.content_type
    )).values():
        if len({record.title for record in group}) > 1:
            add_review(analysis, "title_spacing_variant", group,
                       "Названия различаются пробелами. Автоматического объединения нет.",
                       "Официальное название и ISBN/ID книги.")

    valid_records = [record for record in records if valid_url(record.url)]
    for group in buckets(valid_records, lambda record: record.url_group_id).values():
        if len({record.record_group_id for record in group}) > 1:
            add_review(analysis, "url_record_conflict", group,
                       "Один нормализованный URL связан с разными исходными записями. "
                       "Возможны варианты оформления, конфликт метаданных или сборник.",
                       "Подтверждение метаданных, содержимое страницы и каталожные идентификаторы.")

    mirror_groups = defaultdict(list)
    for record in valid_records:
        parts = urlsplit(record.url)
        base_host = re.sub(r"^s\d+\.", "", parts.hostname)
        mirror_groups[(parts.scheme, base_host, parts.port, parts.path, parts.query, parts.fragment)].append(record)
    for group in mirror_groups.values():
        if len({urlsplit(record.url).hostname for record in group}) > 1:
            add_review(analysis, "possible_mirror", group,
                       "Одинаковый путь на основном хосте и поддомене sN. URL сохранены раздельно.",
                       "Подтверждение зеркала: редирект, canonical или сравнение содержимого.")


def process_workbook(source):
    source = Path(source).resolve()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    records = []
    try:
        sheet = workbook.worksheets[0]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(max_row=1, max_col=6)))
        if headers not in {HEADERS, ("Праообладатель",) + HEADERS[1:]}:
            raise ValueError(f"Неожиданные заголовки первого листа: {headers!r}")
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, max_col=6, values_only=True), 2):
            if all(value is None for value in values):
                continue
            record = Record(source.name, sheet.title, row_number, tuple(values),
                            normalize_identity(values[1]), normalize_identity(values[2]),
                            normalize_url(values[4]), normalize_text(values[3]))
            record.record_group_id = group_id("record", values)
            record.url_group_id = group_id("url", record.url if valid_url(record.url)
                                           else (sheet.title, row_number, record.url))
            book_key = (record.author, record.title, record.content_type)
            if not all(book_key):
                book_key += (sheet.title, row_number)
            record.book_id = group_id("book", book_key)
            records.append(record)
    finally:
        workbook.close()
    if not records:
        raise ValueError("На первом листе нет записей.")
    analysis = Analysis(source, records)
    for duplicate_type, key in (("full_record", "record_group_id"), ("repeated_url", "url_group_id")):
        for identifier, group in buckets(records, lambda record: getattr(record, key)).items():
            if len(group) > 1:
                analysis.duplicates.append({
                    "type": duplicate_type, "group_id": identifier,
                    "source_rows": [record.source_row for record in group],
                    "row_count": len(group), "url_count": len({record.url for record in group}),
                    "book_ids": sorted({record.book_id for record in group}),
                    "urls": sorted({record.url for record in group if record.url}),
                })
                for record in group:
                    record.issues.add(duplicate_type)
    for identifier, group in buckets(records, lambda record: record.book_id).items():
        urls = sorted({record.url for record in group if record.url})
        book = {
            "book_id": identifier, "author": group[0].author, "title": group[0].title,
            "content_type": group[0].content_type,
            "source_rows": [record.source_row for record in group],
            "row_count": len(group), "url_count": len(urls), "urls": urls,
            "original_authors": sorted({str(record.raw[1]) for record in group}),
            "original_titles": sorted({str(record.raw[2]) for record in group}),
        }
        analysis.books.append(book)
        if len(urls) > 1:
            analysis.duplicates.append({
                "type": "book_multiple_urls", "group_id": identifier,
                "source_rows": book["source_rows"], "row_count": len(group),
                "url_count": len(urls), "book_ids": [identifier], "urls": urls,
            })
    find_review_cases(analysis)
    return analysis


def append_values(sheet, values):
    if any(isinstance(value, str) and len(value) > 32767 for value in values):
        raise ValueError("Значение превышает лимит ячейки Excel (32767 символов).")
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # Preserve literal strings, including a leading '='.


def lines(values):
    return "\n".join(map(str, values))


def write_workbook(analysis, output):
    output = Path(output).resolve()
    if output == analysis.source_path:
        raise ValueError("Нельзя перезаписывать исходный Excel.")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    normalized = workbook.create_sheet("Normalized")
    append_values(normalized, HEADERS + (
        "Source row", "Source sheet", "Source file", "Author normalized", "Title normalized",
        "URL normalized", "Content type normalized", "Record group ID", "URL group ID",
        "Book candidate ID", "Issues", "Review IDs",
    ))
    for record in analysis.records:
        append_values(normalized, record.raw + (
            record.source_row, record.source_sheet, record.source_file,
            record.author, record.title, record.url, record.content_type,
            record.record_group_id, record.url_group_id, record.book_id,
            lines(sorted(record.issues)), lines(sorted(record.review_ids)),
        ))
    duplicates = workbook.create_sheet("Duplicates")
    append_values(duplicates, ("Duplicate type", "Group ID", "Source rows", "Row count",
                               "Unique URL count", "Book candidate IDs", "URLs", "Source sheet"))
    for item in analysis.duplicates:
        append_values(duplicates, (item["type"], item["group_id"], lines(item["source_rows"]),
                                   item["row_count"], item["url_count"], lines(item["book_ids"]),
                                   lines(item["urls"]), analysis.records[0].source_sheet))
    books = workbook.create_sheet("Books")
    append_values(books, ("Book candidate ID", "Author normalized", "Title normalized", "Content type",
                         "Source rows", "Row count", "Unique URL count", "URLs", "Original authors",
                         "Original titles", "Candidate state", "Review IDs", "Source sheet"))
    for item in analysis.books:
        reviews = sorted({review_id for record in analysis.records if record.book_id == item["book_id"]
                          for review_id in record.review_ids})
        append_values(books, (item["book_id"], item["author"], item["title"], item["content_type"],
                             lines(item["source_rows"]), item["row_count"], item["url_count"],
                             lines(item["urls"]), lines(item["original_authors"]),
                             lines(item["original_titles"]),
                             "needs_review" if reviews else "candidate_not_verified", lines(reviews),
                             analysis.records[0].source_sheet))
    review = workbook.create_sheet("Review")
    append_values(review, ("Review ID", "Issue code", "Source rows", "Reason", "Needed data",
                          "Book candidate IDs", "Source sheet"))
    for item in analysis.review:
        append_values(review, (item["review_id"], item["code"], lines(item["source_rows"]),
                               item["reason"], item["needed_data"], lines(item["book_ids"]),
                               analysis.records[0].source_sheet))
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.zoomScale = 85
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column in sheet.columns:
            name = str(column[0].value).lower()
            width = 65 if any(word in name for word in ("url", "title", "название", "reason", "needed")) else 27
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet.title == "Normalized":
            sheet.column_dimensions["G"].width = 13
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output)
    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description="Часть 1: локальная обработка первого листа Excel.")
    parser.add_argument("--input", type=Path, help="Исходный XLSX (по умолчанию единственный XLSX в data).")
    parser.add_argument("--output", type=Path, default=Path("results/part1_processed.xlsx"))
    args = parser.parse_args()
    if args.input is None:
        candidates = sorted(path for path in Path("data").glob("*.xlsx") if not path.name.startswith("~$"))
        if len(candidates) != 1:
            parser.error("В data должен быть один XLSX; иначе укажите --input.")
        args.input = candidates[0]
    try:
        analysis = process_workbook(args.input)
        write_workbook(analysis, args.output)
    except (ValueError, OSError) as error:
        parser.error(str(error))
    full = [item for item in analysis.duplicates if item["type"] == "full_record"]
    print(f"Records: {len(analysis.records)}")
    print(f"Unique URLs: {len({record.url for record in analysis.records})}")
    print(f"Repeated records: {sum(item['row_count'] - 1 for item in full)}")
    print(f"Book candidates: {len(analysis.books)}")
    print(f"Review cases: {len(analysis.review)}")
    print(f"Output: {args.output}")


if __name__ == "__main__":
    main()
