"""Part 2: catalogue matching, explainable decisions, CLI and compact exports."""

import argparse
from contextlib import nullcontext
from collections import Counter
from dataclasses import dataclass, replace
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import time

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import requests

from pages import (authors_match, extract_page, fetch_url, number_value, pdf_identity_matches,
                   text_key, title_byline_role, title_key, title_punctuation_conflict, volume_conflict, volume_markers)
from part1 import group_id, normalize_identity, normalize_text, normalize_url


@dataclass(frozen=True)
class Book:
    id: str
    author: str
    title: str
    content_type: str = "ebook"


def author_matches(expected, observed):
    return authors_match([expected], [observed])


def title_relation(book, observed):
    expected = title_key(book.title)
    candidate = title_key(observed)
    family = lambda title: re.split(r"\b(?:том|книга|часть|ч|volume|vol|book|part)\b", title_key(title))[0].strip()
    conflicting_volumes = (volume_conflict(book.title, observed)
                           and family(book.title) == family(observed))
    if "|" in observed and conflicting_volumes:
        return "volume_conflict"
    if candidate == expected:
        return "punctuation_variant" if title_punctuation_conflict(book.title, observed) else "exact"
    if title_byline_role(observed, book.title, [book.author]) == "author_list":
        return "exact"
    # Only known service words around a full title are accepted as decoration.
    if f" {expected} " in f" {candidate} ":
        rest = (" " + candidate + " ").replace(" " + expected + " ", " ").strip()
        if re.fullmatch(r"(?:(?:читать|книгу|скачать|онлайн|полностью|бесплатно|аудиокнига|аудиокнигу|fb2|epub|pdf|txt|mp3|торрент)\s*)+", rest):
            return "punctuation_variant" if title_punctuation_conflict(book.title, observed) else "exact"
        return "punctuation_variant" if "|" in observed or "|" in book.title else "partial"
    if conflicting_volumes:
        return "volume_conflict"
    if "ё" in expected or "ё" in candidate:
        folded = replace(book, title=normalize_text(book.title), author=normalize_text(book.author))
        if title_relation(folded, normalize_text(observed)) in {"exact", "partial"}:
            return "yo_variant"
    return "none"


def contains_title(text, title):
    return bool(title_key(title)) and f" {title_key(title)} " in f" {title_key(text)} "


def match_one(page, book):
    subjects = list(page.subjects)
    if (page.kind == "pdf" and "pdf_identity_unresolved" not in page.issues
        and (not page.title or not page.authors)
        and pdf_identity_matches(page.opening, book.title, [book.author])):
        # Only a title/byline prefix can identify the document, never arbitrary mentions.
        subjects.append({"title": book.title, "authors": [book.author], "source": "pdf_opening",
                         "scope": "pdf", "signals": [signal for signal in page.signals
                                                     if signal.get("binding", "bound") == "bound"]})
    outcomes, variants = [], []
    for subject in subjects:
        relation = title_relation(book, subject["title"])
        if relation == "none":
            continue
        yo_variant = relation == "yo_variant"
        punctuation_variant = relation == "punctuation_variant"
        if relation == "volume_conflict":
            state = "volume_conflict"
        elif relation in {"partial", "yo_variant", "punctuation_variant"}:
            state = "ambiguous"
        elif not subject["authors"]:
            state = "author_missing"
        elif not authors_match([book.author], subject["authors"]):
            yo_variant = authors_match([book.author.replace("ё", "е").replace("Ё", "Е")],
                                       [a.replace("ё", "е").replace("Ё", "Е") for a in subject["authors"]])
            state = "ambiguous" if yo_variant else "author_conflict"
        else:
            state = "matched"
        outcome = {"state": state, "book_id": book.id, "subject": subject,
                          "signals": subject["signals"],
                          "unresolved_signals": subject.get("unresolved_signals", []),
                          "detected_formats": subject.get("detected_formats", sorted({
                              signal["format"] for signal in subject["signals"]})),
                         "identity_evidence": [f"{subject['source']}: {subject['title'][:180]}",
                                               "Автор: " + "; ".join(subject["authors"])[:180]]}
        if yo_variant or punctuation_variant:
            reason = ("Различие е/ё требует подтверждения идентичности; поискового сходства недостаточно."
                      if yo_variant else "Пунктуация названия различается; одинаковых слов недостаточно для подтверждения идентичности.")
            outcome["unresolved_signals"] = [
                {**signal, "binding": "unresolved", "binding_reason": reason}
                for signal in subject["signals"] + subject.get("unresolved_signals", [])]
            outcome["signals"] = []
            outcome["identity_evidence"].append(reason + " Каталог: " + book.title + "; " + book.author)
            variants.append(outcome)
        outcomes.append(outcome)
    priority = {"matched": 0, "author_conflict": 1, "author_missing": 2,
                "volume_conflict": 3, "ambiguous": 4}
    if outcomes:
        outcomes.sort(key=lambda item: (priority[item["state"]], not bool(item["signals"])))
        chosen = outcomes[0]
        if chosen["state"] == "matched" and any(
            item["state"] == "author_conflict" and (
                item["subject"].get("scope_id") == chosen["subject"].get("scope_id")
                and (item["subject"]["source"] != "jsonld_book" or item["subject"].get("primary_metadata"))
            ) for item in outcomes
        ):
            chosen = {**chosen, "state": "author_conflict"}
        if chosen["state"] == "matched" and not chosen["signals"]:
            chosen = next((item for item in variants if item["unresolved_signals"]), chosen)
        return chosen
    if contains_title(page.text, book.title) or contains_title(normalize_text(page.text), normalize_text(book.title)):
        return {"state": "mentioned", "book_id": book.id, "signals": [],
                "identity_evidence": ["Название найдено в тексте, но не в подтверждённой карточке книги."]}
    return {"state": "no_match", "book_id": None, "signals": [], "identity_evidence": []}


def match_book(page, catalog, expected_id=None):
    catalog = list(catalog)
    outcomes = {book.id: match_one(page, book) for book in catalog}
    for subject in page.subjects:
        if subject["source"] not in {"primary", "book_card"} or not any(
            title_key(book.title) == title_key(subject["title"])
            and not title_punctuation_conflict(book.title, subject["title"])
            and authors_match([book.author], subject["authors"]) for book in catalog
        ):
            continue
        metadata = next((item for item in page.subjects if item["source"] == "jsonld_book"
                         and item.get("primary_metadata") and item.get("scope_id") == subject.get("scope_id")
                         and authors_match(subject["authors"], item["authors"])
                         and title_byline_role(subject["title"], item["title"], subject["authors"])), None)
        if metadata is None:
            continue
        reason = ("Полное каталожное заглавие совпало с HTML, но противоречит JSON-LD; роль подписи автора неоднозначна: "
                  + subject["title"] + " / " + metadata["title"])
        for outcome in outcomes.values():
            if outcome["state"] == "matched" and outcome["subject"].get("scope_id") == subject.get("scope_id"):
                outcome["state"] = "ambiguous"
                outcome["identity_evidence"].append(reason)
                outcome["unresolved_signals"] = [{**signal, "binding": "unresolved", "binding_reason": reason}
                                                 for signal in outcome["signals"] + outcome["unresolved_signals"]]
                outcome["signals"] = []
    if expected_id is not None:
        selected = outcomes.get(expected_id, {"state": "catalog_missing", "book_id": expected_id,
                                              "signals": [], "identity_evidence": []})
    else:
        hits = [item for item in outcomes.values() if item["state"] == "matched"]
        if len(hits) == 1:
            selected = hits[0]
        elif len(hits) > 1:
            selected = {"state": "ambiguous", "book_id": None, "signals": [],
                        "identity_evidence": ["Совпало несколько книг каталога: " + ", ".join(item["book_id"] for item in hits)]}
        else:
            selected = next((item for item in outcomes.values() if item["state"] != "no_match"),
                            {"state": "no_match", "book_id": None, "signals": [], "identity_evidence": []})
    return {**selected, "other_catalog_matches": [identifier for identifier, item in outcomes.items()
                                                  if item["state"] == "matched" and identifier != selected["book_id"]]}


REASONS = {
    "author_conflict": "Название похоже, но извлечённый автор противоречит каталогу.",
    "author_missing": "Название найдено, но автор страницы не установлен; одного названия недостаточно.",
    "volume_conflict": "У совпадающей серии различаются номера тома/части/книги.",
    "ambiguous": "Неполное или неоднозначное совпадение названия/кандидатов каталога.",
    "catalog_missing": "Ожидаемая книга отсутствует в переданном каталоге.",
    "collection_unresolved": "Книга упомянута в сборнике, но контент/файл не привязан к её отдельной карточке.",
    "mentioned": "Обнаружено только упоминание книги без уверенной идентификации карточки.",
    "format_conflict": "Произведение совпало, но признаки контента относятся к другому формату (audio/ebook).",
    "no_match": "Страница проанализирована; уверенного совпадения с проверяемой книгой не найдено.",
    "content_not_confirmed": "Название и автор совпали, но конкретные признаки размещения контента не найдены.",
    "loading_shell": "Получена загрузочная оболочка: содержимое книги ещё не доступно для анализа.",
    "content_requires_js": "Заявлены чтение/плеер, но связанное содержимое ещё не получено; требуется браузерная проверка.",
    "pdf_preview_incomplete": "Ограниченного чтения PDF недостаточно для надёжного вывода; полная книга не скачивается.",
    "service_error": "Сайт вернул служебную страницу ошибки, а не содержимое книги.",
    "content_scope_unresolved": "Книга совпала, но принадлежность файла, текста или плеера этой книге не установлена.",
    "placement_signs": "Название и автор совпали; найдены связанные с книгой конкретные признаки контента. "
                       "Это не проверка полноты файла и не вывод о нарушении авторских прав.",
    "access_challenge": "CAPTCHA/проверка браузера или страница ограничения доступа; обход не выполнялся.",
    "login_required": "Для просмотра требуется авторизация; вход не выполнялся.",
    "insufficient_html": "HTML не содержит достаточно доступных данных (возможно, требуется JavaScript).",
    "metadata_title_conflict": "Заголовок страницы и структурированные метаданные книги противоречат друг другу.",
    "metadata_author_conflict": "Автор основной книги в JSON-LD противоречит автору в DOM.",
    "visibility_unverified": "Принадлежность контента видимой области требует проверки вычисленных CSS-стилей в браузере.",
    "pdf_no_text": "PDF не содержит достаточного извлекаемого текста; OCR не выполнялся.",
    "pdf_identity_unresolved": "Название и автор из метаданных PDF не подтверждены началом текста; "
                               "возможен конфликт метаданных и содержимого, идентичность книги не установлена.",
    "pdf_document_role_unresolved": "PDF идентифицирует книгу, но имеются признаки пересказа, рецензии, критического разбора или учебных материалов; текст произведения не подтверждён.",
    "encrypted_pdf": "PDF зашифрован; содержимое недоступно без пароля.",
    "extraction_error": "Документ получен, но не удалось надёжно извлечь данные.",
    "unsupported_content_type": "Получен неподдерживаемый тип ответа, а не HTML/PDF.",
    "partial_document": "Извлечение ограничено частью документа; отрицательный вывод ненадёжен.",
    "not_loaded": "Ответ пуст: извлекаемые данные страницы отсутствуют.",
    "catalog_format_unknown": "Тип контента каталога не поддержан: ожидается ebook или audio/audiobook.",
    "different_pdf_book": "PDF идентифицирован по метаданным и началу текста как другая книга; "
                          "совпадение с ожидаемой книгой не подтверждено (прочитан ограниченный фрагмент).",
}


def classify(document, page, match, catalog):
    if document.http_status in {404, 410} and not document.error_code.startswith("browser_"):
        return 0, f"http_{document.http_status}", "Страница отсутствует (HTTP 404/410); размещение по этому URL не подтверждено."
    if document.error_code:
        return 2, document.error_code, f"Загрузка не завершена: {document.error_code}. {document.error_detail}".strip()
    confirmed = match["state"] == "matched" and bool(match["signals"])
    weak = {"insufficient_html", "access_challenge", "login_required", "content_requires_js"}
    for issue in page.issues:
        if not page.issue_applies(issue, match.get("subject")):
            continue
        if issue in weak and issue not in page.blocking_issues:
            if confirmed:
                continue
            if issue == "content_requires_js" and not page.dynamic_pending and match.get("subject") and not match["subject"].get("requires_js"):
                continue
        if issue in REASONS:
            return 2, issue, REASONS[issue]
    state = match["state"]
    if page.is_collection and state in {"mentioned", "author_missing", "ambiguous"}:
        return 2, "collection_unresolved", REASONS["collection_unresolved"]
    if state == "matched":
        if page.is_collection and not match["signals"]:
            return 2, "collection_unresolved", REASONS["collection_unresolved"]
        book = next(book for book in catalog if book.id == match["book_id"])
        supplied_kind = normalize_text(book.content_type)
        if supplied_kind not in {"ebook", "audio", "audiobook", "аудиокнига"}:
            return 2, "catalog_format_unknown", REASONS["catalog_format_unknown"]
        if match["signals"]:
            return 1, "placement_signs", REASONS["placement_signs"]
        if match.get("unresolved_signals"):
            return 2, "content_scope_unresolved", REASONS["content_scope_unresolved"]
        if page.partial:
            return 2, "partial_document", REASONS["partial_document"]
        # PDF thresholds establish positive evidence, not the absence of the work.
        return (2 if page.kind == "pdf" else 0), "content_not_confirmed", REASONS["content_not_confirmed"]
    if state != "no_match":
        return 2, state, REASONS[state]
    if page.partial:
        return 2, "partial_document", REASONS["partial_document"]
    if page.kind == "pdf" and page.title and page.authors and contains_title(page.opening, page.title):
        return 0, "different_pdf_book", REASONS["different_pdf_book"]
    return 0, "no_match", REASONS["no_match"]


def evaluate_page(document, catalog, expected_id=None):
    # A failed browser wait can still provide observed metadata. Keep those observations,
    # while classification uses the original failure and therefore cannot become positive.
    observed = replace(document, error_code="") if document.body and document.error_code.startswith("browser_") else document
    page = extract_page(observed)
    match = match_book(page, catalog, expected_id)
    status, reason_code, reason = classify(document, page, match, catalog)
    evidence = match["identity_evidence"] + [
        f"{signal['detail']} {signal['url']}" + (
            f" Область: {signal['scope_locator']}; заголовок: {signal['scope_title']}."
            if signal.get("binding") == "bound" else "") for signal in match["signals"]]
    if reason_code == "content_scope_unresolved":
        evidence += [f"Не привязано: {signal['detail']} {signal['binding_reason']}"[:400]
                     for signal in match.get("unresolved_signals", [])[:3]]
    if reason_code == "pdf_identity_unresolved":
        evidence.append("Начало текста PDF: " + page.summary()["text_excerpt"])
    if reason_code == "pdf_document_role_unresolved":
        evidence.append("Признаки роли PDF: " + page.metadata.get("document_role_evidence", ""))
    if reason_code in {"metadata_title_conflict", "metadata_author_conflict"}:
        scopes = page.issue_scopes.get(reason_code, [])
        evidence += [f"JSON-LD: {scope['metadata_title']}; автор: {'; '.join(scope['metadata_authors'])}; "
                     f"область: {scope['title']} {scope['scope_locator']}"[:400] for scope in scopes]
        if not scopes:
            evidence += [f"Основной JSON-LD: {subject['title']}; автор: {'; '.join(subject['authors'])}"[:400]
                         for subject in page.subjects if subject.get("primary_metadata")]
    for issue, scopes in page.issue_scopes.items():
        if not page.issue_applies(issue, match.get("subject")):
            evidence.insert(0, "Ошибка независимой карточки не относится к выбранной книге: " + issue + "; " +
                            "; ".join(f"{scope['title']} {scope['scope_locator']} (JSON-LD: {scope['metadata_title']})"
                                      for scope in scopes)[:300])
    if not evidence and page.title:
        evidence = ["Извлечённый заголовок: " + page.title[:200],
                    "Извлечённые авторы: " + "; ".join(page.authors)[:180]]
    if reason_code == "partial_document" and page.kind == "pdf":
        evidence.append(f"Прочитано страниц PDF: {page.pdf_pages_read}/{page.pdf_page_count}; "
                        "анализ документа неполон, отсутствие произведения не установлено.")
    if document.error_code:
        evidence = [f"HTTP: {document.http_status}; {document.error_code}; {document.error_detail}"[:500]]
    public_match = {key: value for key, value in match.items() if key not in {"subject", "signals"}}
    if document.error_code or (status == 2 and any(issue in page.issues for issue in {
        "not_loaded", "access_challenge", "login_required", "insufficient_html", "pdf_no_text",
        "encrypted_pdf", "extraction_error", "unsupported_content_type", "loading_shell", "content_requires_js",
        "pdf_preview_incomplete", "service_error",
    })):
        public_match["state"] = "unavailable"
    if "pdf_identity_unresolved" in page.issues:
        public_match["state"] = "ambiguous"
    result = {
        "url": document.url, "final_url": document.final_url, "http_status": document.http_status,
        "checked_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "expected_book_id": expected_id, "fetch_error": document.error_code,
        "fetch_error_detail": document.error_detail,
        "redirects": document.redirects, "response_bytes": document.size_bytes or len(document.body),
        "elapsed_seconds": document.elapsed_seconds, "extracted": page.summary(),
        "match": public_match, "content_signals": match["signals"],
        "status": status, "reason_code": reason_code, "reason": reason, "evidence": list(dict.fromkeys(evidence))[:8],
    }
    book = next((book for book in catalog if book.id == match["book_id"]), None)
    formats = sorted({signal["format"] for signal in match["signals"] if signal.get("format") != "unknown"})
    result["catalog_format"] = book.content_type if book else None
    result["observed_formats"] = formats
    expected_format = "audio" if book and normalize_text(book.content_type) in {"audio", "audiobook", "аудиокнига"} else "ebook"
    result["format_difference"] = bool(formats and expected_format not in formats) if book else None
    return complete_evidence(result)


def analyze_url(target, catalog, fetcher=fetch_url, browser=None):
    """HTTP first, optional bounded browser fallback, the same extract/match/classify pipeline."""
    initial = fetcher(target["url"])
    result = evaluate_page(initial, catalog, target.get("expected_book_id"))
    http_attempt = {"http_status": initial.http_status, "final_url": initial.final_url,
                    "error": initial.error_code, "status": result["status"], "reason_code": result["reason_code"]}
    result.update(method="HTTP", initial_http_status=initial.http_status, http_attempt=http_attempt,
                  browser_attempt={"attempted": False, "ready": False, "reason": "Не требуется или отключён."})
    eligible = result["reason_code"] in {"loading_shell", "insufficient_html", "access_challenge", "not_loaded",
                                          "content_requires_js", "visibility_unverified", "http_403", "http_503"}
    if eligible and browser is not None and result["extracted"]["kind"] != "pdf":
        from browser import BrowserProbe

        def target_ready(page):
            selected = match_book(page, catalog, target.get("expected_book_id"))
            if selected["state"] == "matched" and selected["signals"]:
                return True
            if page.dynamic_pending:
                return False
            subject = selected.get("subject", {})
            if subject:
                return not subject.get("requires_js")
            # A fully identified different book can still establish a negative result.
            return selected["state"] == "no_match" and any(
                item["authors"] and item["signals"] for item in page.subjects)

        if isinstance(browser, BrowserProbe):
            rendered = browser.check(target["url"], ready=target_ready,
                                     html_document=result["extracted"]["kind"] == "html")
        else:
            rendered = browser.check(target["url"])
        result = evaluate_page(rendered, catalog, target.get("expected_book_id"))
        result.update(method="Chromium", initial_http_status=initial.http_status, http_attempt=http_attempt,
                      browser_attempt={"attempted": True, "ready": not bool(rendered.error_code),
                                       "error": rendered.error_code, "detail": rendered.error_detail,
                                       "http_status": rendered.http_status, "final_url": rendered.final_url,
                                       "version": getattr(browser, "version", "")})
    result.update({key: value for key, value in target.items() if key != "url"})
    return result


def complete_evidence(result):
    """Compact proof also for empty/blocked pages; uses only recorded summaries."""
    if result["reason_code"] == "format_conflict":
        proof = "Формат книги: " + ", ".join(result["match"].get("detected_formats", []))
        if proof not in result["evidence"]:
            result["evidence"].append(proof)
    if not result["evidence"]:
        result["evidence"] = [
            f"HTTP {result['http_status']}; ответ {result['response_bytes']} байт; "
            f"извлечено {result['extracted']['text_chars']} символов.",
            "Проблемы извлечения: " + ", ".join(result["extracted"]["issues"]),
        ]
    result["evidence"] = list(dict.fromkeys(result["evidence"]))
    return result


def load_inputs(path):
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    books, targets, identities = {}, [], {}
    try:
        sheet = workbook.worksheets[1]
        for number, values in enumerate(sheet.iter_rows(min_row=2, max_col=6, values_only=True), 2):
            if all(value is None for value in values):
                continue
            _, author, title, content_type, url, _ = values
            if not author or not title:
                raise ValueError(f"Строка {number}: в мини-каталоге отсутствует автор/название.")
            key = (normalize_identity(author), normalize_identity(title), normalize_text(content_type))
            identifier = group_id("book", key)
            if identifier in identities and identities[identifier] != key:
                raise ValueError(f"Коллизия идентификатора книги {identifier}, строка {number}; кандидаты не перезаписаны.")
            identities[identifier] = key
            books.setdefault(identifier, Book(identifier, str(author), str(title), str(content_type or "ebook")))
            targets.append({"url": normalize_url(url), "expected_book_id": identifier,
                            "source_file": path.name, "source_sheet": sheet.title, "source_row": number,
                            "expected_author": author, "expected_title": title, "expected_type": content_type})
    finally:
        workbook.close()
    return list(books.values()), targets


def load_json_catalog(path):
    values = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if not isinstance(values, list) or not values:
        raise ValueError("Каталог JSON должен быть непустым списком книг.")
    catalog = []
    for item in values:
        if not isinstance(item, dict) or not all(
            isinstance(item.get(key), str) and item[key].strip() for key in ("author", "title")
        ) or not isinstance(item.get("content_type", "ebook"), str):
            raise ValueError("Каждая книга должна иметь непустые author и title.")
        identifier = str(item.get("id") or group_id("book", (item["author"], item["title"], item.get("content_type", "ebook"))))
        catalog.append(Book(identifier, item["author"], item["title"], item.get("content_type", "ebook")))
    if len({book.id for book in catalog}) != len(catalog):
        raise ValueError("Идентификаторы книг каталога должны быть уникальны.")
    return catalog


def write_results(results, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    with (output_dir / "part2_results.jsonl").open("w", encoding="utf-8", newline="\n") as stream:
        for result in results:
            stream.write(json.dumps(result, ensure_ascii=False) + "\n")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    fields = ["source_file", "source_sheet", "source_row", "url", "expected_author", "expected_title",
              "expected_type", "expected_book_id", "http_status", "final_url", "status", "reason_code", "reason",
              "method", "initial_http_status", "http_attempt", "browser_attempt", "catalog_format",
              "observed_formats", "format_difference",
              "match_state", "matched_book_id", "title", "h1", "authors", "metadata",
              "document_type", "pdf_page_count", "pdf_pages_read", "partial", "catalog_match_details",
              "text_chars", "text_excerpt", "content_signals", "evidence", "extraction_issues",
              "observed_content_signals", "interaction_hints", "detected_formats",
              "fetch_error_detail", "redirects", "response_bytes", "elapsed_seconds", "checked_at"]
    sheet.append(fields)
    for result in results:
        flattened = {**result, "match_state": result["match"]["state"],
                     "matched_book_id": result["match"]["book_id"], **result["extracted"],
                     "catalog_match_details": result["match"], "document_type": result["extracted"]["kind"],
                     "extraction_issues": result["extracted"]["issues"]}
        values = [flattened.get(key) for key in fields]
        values = [json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value for value in values]
        if any(isinstance(value, str) and len(value) > 32767 for value in values):
            raise ValueError("Результат превышает лимит ячейки Excel.")
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for column, name in enumerate(fields, 1):
        sheet.column_dimensions[get_column_letter(column)].width = 55 if name in {
            "url", "final_url", "reason", "evidence", "content_signals", "expected_title", "text_excerpt"
        } else 25
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output_dir / "part2_results.xlsx")
    workbook.close()


def main():
    parser = argparse.ArgumentParser(description="Часть 2: ограниченная HTTP-проверка URL и сопоставление с каталогом.")
    parser.add_argument("--input", type=Path, help="XLSX: второй лист — мини-каталог и URL.")
    parser.add_argument("--catalog", type=Path, help="Другой каталог: JSON со списком книг.")
    parser.add_argument("--urls", type=Path, help="Другой список: UTF-8, один URL на строку.")
    parser.add_argument("--output-dir", type=Path, default=Path("results"))
    parser.add_argument("--delay", type=float, default=1, help="Пауза между URL, секунды (по умолчанию 1).")
    parser.add_argument("--browser", choices=("auto", "off"), default="auto", help="Fallback Chromium (auto) или только HTTP (off).")
    parser.add_argument("--browser-channel", choices=("auto", "chromium", "msedge", "chrome"), default="auto")
    parser.add_argument("--browser-executable", type=Path, help="Путь к уже установленному Chromium (необязательно).")
    parser.add_argument("--browser-timeout", type=int, default=25, help="Бюджет браузерного ожидания, секунды.")
    args = parser.parse_args()
    if bool(args.catalog) != bool(args.urls) or args.input and args.catalog:
        parser.error("Используйте --input либо пару --catalog и --urls.")
    if args.delay < 0 or args.browser_timeout <= 0:
        parser.error("--delay должен быть >=0, --browser-timeout >0.")
    try:
        if args.catalog:
            catalog = load_json_catalog(args.catalog)
            targets = [{"url": normalize_url(line), "expected_book_id": None,
                        "source_file": args.urls.name, "source_sheet": None, "source_row": number}
                       for number, line in enumerate(args.urls.read_text(encoding="utf-8-sig").splitlines(), 1)
                       if line.strip()]
        else:
            if args.input is None:
                candidates = [path for path in Path("data").glob("*.xlsx") if not path.name.startswith("~$")]
                if len(candidates) != 1:
                    parser.error("Укажите --input: в data нет единственного XLSX.")
                args.input = candidates[0]
            catalog, targets = load_inputs(args.input)
        if not targets:
            raise ValueError("Список URL пуст.")
    except (ValueError, OSError, IndexError) as error:
        parser.error(str(error))
    results, cache = [], {}
    if args.browser == "auto":
        from browser import BrowserProbe
        browser_context = BrowserProbe(timeout_ms=args.browser_timeout * 1000,
                                       channel=args.browser_channel, executable=args.browser_executable)
    else:
        browser_context = nullcontext(None)
    with requests.Session() as session, browser_context as browser:
        def fetch_cached(url):
            if url not in cache:
                if cache:
                    time.sleep(args.delay)
                cache[url] = fetch_url(url, session=session)
            return cache[url]

        for target in targets:
            result = analyze_url(target, catalog, fetcher=fetch_cached, browser=browser)
            results.append(result)
            print(f"Row {target['source_row']}: {result['method']} HTTP {result['http_status']} -> {result['status']} ({result['reason_code']})", flush=True)
    write_results(results, args.output_dir)
    print(f"Results: {len(results)}; statuses: {dict(sorted(Counter(item['status'] for item in results).items()))}")
    print(f"Output: {args.output_dir / 'part2_results.xlsx'}; {args.output_dir / 'part2_results.jsonl'}")


if __name__ == "__main__":
    main()
