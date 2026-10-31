"""RT-04: compare complete coauthor sets, never a count/character prefix."""

from io import BytesIO
import json
import sys

import openpyxl
import pytest
import requests
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from browser import BrowserProbe
from pages import FetchResult, Limits, authors_match, extract_page, fetch_url
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main


TITLE = "Silent Harbor"
URL = "https://offline.authors/book"
AUTHORS = [f"Writer Number{i}" for i in range(11)]
FILE = '<a href="/files/123.epub">Download EPUB</a>'
DETAILS = '<p>Public bibliographic information and available reading formats.</p>'
FORMS = ["dom", "packed", "labelled", "labelled-packed", "labelled-split", "microdata", "author-links", "jsonld", "ids"]


@pytest.fixture(autouse=True)
def no_live_http(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden in author completeness regressions")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def book(names, identifier="target", title=TITLE):
    return Book(identifier, "; ".join(names), title)


def byline(names, form):
    packed = "; ".join(names)
    if form == "packed":
        return f'<p class="book-author">{packed}</p>'
    if form == "labelled":
        return "".join(f"<p>Author: {name}</p>" for name in names)
    if form == "labelled-packed":
        return f"<p>Author: {packed}</p>"
    if form == "labelled-split":
        return f"<p><b>Author:</b><span>{packed}</span></p>"
    if form == "microdata":
        return "".join(f'<span itemprop="author">{name}</span>' for name in names)
    if form == "author-links":
        return "".join(f'<a href="/authors/{i}" rel="author">{name}</a>' for i, name in enumerate(names))
    return "".join(f'<span class="book-author">{name}</span>' for name in names)


def schema(names, linked=False, title=TITLE, prefix="own"):
    data = {"@context": "https://schema.org", "@type": "https://schema.org/Book", "@id": "#" + prefix,
            "name": title, "author": [{"@id": f"#{prefix}-{i}"} if linked else {"@type": "Person", "name": name}
                                       for i, name in enumerate(names)]}
    if linked:
        data = {"@graph": [{"@type": "WebPage", "@id": URL, "mainEntity": {"@id": "#" + prefix}}, data,
                            *[{"@type": "Person", "@id": f"#{prefix}-{i}", "name": name} for i, name in enumerate(names)]]}
    return '<script type="application/ld+json">' + json.dumps(data) + "</script>"


def card(names, form="dom", title=TITLE, content=FILE, metadata_names=None):
    structured = schema(names if metadata_names is None else metadata_names, form == "ids", title,
                        "own" if title == TITLE else "other") if form in {"jsonld", "ids"} else ""
    return f"<article><h1>{title}</h1>" + byline(names, "dom" if structured else form) + DETAILS + structured + content + "</article>"


def document(names, form="dom", body=None, head="", content=FILE):
    html = '<html><head><meta charset="utf-8">' + head + '</head><body><main>' + (card(names, form, content=content) if body is None else body) + "</main></body></html>"
    return FetchResult(URL, URL, 200, {"Content-Type": "text/html; charset=utf-8"}, html.encode())


@pytest.mark.parametrize("count", [10, 11])
@pytest.mark.parametrize("form", FORMS)
def test_complete_sets_match_in_each_supported_representation(count, form):
    names = AUTHORS[:count]
    expected = book(names)
    initial = document(names, form)
    result = evaluate_page(initial, [expected], expected.id)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")
    # Export summaries must not silently present a complete list as only 10 authors.
    assert authors_match([expected.author], result["extracted"]["authors"])
    assert any(authors_match([expected.author], subject["authors"]) for subject in result["extracted"]["subjects"])
    assert authors_match([expected.author], result["content_signals"][0]["scope_authors"])


@pytest.mark.parametrize("form", ["dom", "packed", "ids"])
@pytest.mark.parametrize(("observed", "status"), [
    (AUTHORS[:10], 2),
    (AUTHORS + ["Extra Writer"], 2),
    (AUTHORS[:10] + ["Different Writer"], 2),
    (list(reversed(AUTHORS)), 1),
    (AUTHORS + [AUTHORS[0], AUTHORS[-1]], 1),
    ([name.replace("Writer Number", "Number") + " Writer" for name in AUTHORS], 1),
], ids=["removed", "added", "replaced", "reordered", "duplicates", "surname-first"])
def test_add_remove_replace_order_and_duplicates_use_complete_sets(form, observed, status):
    expected = book(AUTHORS)
    result = evaluate_page(document(observed, form), [expected], expected.id)
    assert result["status"] == status


@pytest.mark.parametrize("form", FORMS)
def test_eleven_authors_are_not_equal_to_the_first_ten(form):
    expected = book(AUTHORS[:10])
    result = evaluate_page(document(AUTHORS, form), [expected], expected.id)
    assert result["status"] == 2


@pytest.mark.parametrize("count", [25, 50])
def test_no_replacement_arbitrary_author_count_limit(count):
    names = [f"Writer Number{i}" for i in range(count)]
    expected = book(names)
    result = evaluate_page(document(names), [expected], expected.id)
    assert result["status"] == 1
    assert len(result["extracted"]["authors"]) == count
    assert len(result["content_signals"][0]["scope_authors"]) == count


@pytest.mark.parametrize("complete", [False, True])
def test_long_labelled_member_cannot_be_dropped_from_the_set(complete):
    names = AUTHORS[:10] + ["Writer " + "Longname" * 15]
    expected = book(names if complete else names[:10])
    result = evaluate_page(document(names, "labelled"), [expected], expected.id)
    assert result["status"] == (1 if complete else 2)


def padded_author_value():
    prefix = "; ".join(AUTHORS[:10])
    return prefix + " " * (1000 - len(prefix)) + "; " + AUTHORS[-1]


@pytest.mark.parametrize("field", ["author", "book:author"])
@pytest.mark.parametrize("complete", [False, True])
def test_html_metadata_author_is_not_cut_at_1000_characters(field, complete):
    expected = book(AUTHORS if complete else AUTHORS[:10])
    head = f'<meta name="{field}" content="{padded_author_value()}">'
    body = f"<article><h1>{TITLE}</h1>{DETAILS}{FILE}</article>"
    result = evaluate_page(document([], body=body, head=head), [expected], expected.id)
    assert result["status"] == (1 if complete else 2)
    assert result["extracted"]["metadata"][field] == padded_author_value()


def pdf(author):
    writer = PdfWriter()
    writer.add_metadata({"/Title": TITLE, "/Author": author})
    for _ in range(2):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        lines = [TITLE, author, "Chapter 1", *(["The traveler crossed the harbor and returned to the old house before sunset."] * 24)]
        escaped = [line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)') for line in lines]
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 10 Tf 30 760 Td 12 TL ' + ' '.join(f'({line}) Tj T*' for line in escaped) + ' ET').encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return FetchResult(URL, URL, 200, {"Content-Type": "application/pdf"}, output.getvalue())


def test_pdf_metadata_author_is_not_a_truncated_prefix_confirmation():
    expected = book(AUTHORS[:10])
    result = evaluate_page(pdf(padded_author_value()), [expected], expected.id)
    assert result["status"] == 2
    assert result["extracted"]["metadata"]["/Author"] == padded_author_value()


@pytest.mark.parametrize("count", [10, 11])
def test_complete_short_pdf_author_sets_keep_the_positive(count):
    expected = book(AUTHORS[:count])
    assert evaluate_page(pdf(expected.author), [expected], expected.id)["status"] == 1


@pytest.mark.parametrize("complete", [False, True])
def test_neighbouring_bibliography_compares_the_whole_byline(complete):
    expected = book(AUTHORS if complete else AUTHORS[:10])
    content = f'<p>{TITLE} by {"; ".join(AUTHORS)}</p>' + FILE
    result = evaluate_page(document(AUTHORS if complete else AUTHORS[:10], content=content), [expected], expected.id)
    assert result["status"] == (1 if complete else 2)


@pytest.mark.parametrize("fault", [None, "target", "other"])
def test_bm_open_16_many_author_conflicts_stay_local(fault):
    other = book(AUTHORS[:10], "other", "Winter Garden")
    own = book(AUTHORS)
    own_card = card(AUTHORS, "ids", metadata_names=AUTHORS[:10] if fault == "target" else AUTHORS)
    other_card = card(AUTHORS[:10], "jsonld", title=other.title,
                      content='<a href="/files/456.epub">Download EPUB</a>',
                      metadata_names=AUTHORS if fault == "other" else AUTHORS[:10])
    initial = document([], body=own_card + other_card)
    for expected, status in [(own, 2 if fault == "target" else 1), (other, 2 if fault == "other" else 1)]:
        result = evaluate_page(initial, [own, other], expected.id)
        assert result["status"] == status
        if fault == expected.id or (fault == "target" and expected.id == own.id):
            assert result["reason_code"] == "metadata_author_conflict"


@pytest.mark.parametrize("own_file", [False, True])
def test_foreign_card_does_not_supply_a_many_author_books_file(own_file):
    own = book(AUTHORS)
    other = book(AUTHORS[:10], "other", "Winter Garden")
    body = card(AUTHORS, content=FILE if own_file else "") + card(AUTHORS[:10], title=other.title)
    result = evaluate_page(document([], body=body), [own, other], own.id)
    assert result["status"] == (1 if own_file else 0)
    assert all(signal["scope_title"] == TITLE for signal in result["content_signals"])


class Session:
    def __init__(self, body, declared=None):
        self.body, self.declared = body, declared
    def get(self, url, **kwargs):
        response = requests.Response()
        response.status_code = 200
        response.url = url
        response.headers["Content-Type"] = "text/html; charset=utf-8"
        response.headers["Content-Length"] = str(len(self.body) if self.declared is None else self.declared)
        response.raw = BytesIO(self.body)
        return response


@pytest.mark.parametrize("exceeded", [False, True])
def test_document_byte_budget_is_the_protective_boundary(exceeded):
    limit = Limits().max_bytes
    initial = document(AUTHORS)
    body = initial.body + b'<!--' + b' ' * (limit - len(initial.body) - 7) + b'-->'
    if exceeded:
        body += b' '
    expected = book(AUTHORS)
    fetched = fetch_url(URL, session=Session(body))
    result = evaluate_page(fetched, [expected], expected.id)
    assert result["status"] == (2 if exceeded else 1)
    if exceeded:
        assert fetched.error_code == "response_too_large"
        assert fetched.body == b''


@pytest.mark.parametrize("complete", [False, True])
def test_text_excerpt_budget_cannot_hide_late_coauthor(complete):
    expected = book(AUTHORS if complete else AUTHORS[:10])
    body = (f'<article><h1>{TITLE}</h1>' + byline(AUTHORS[:10], "dom")
            + '<p>' + 'x ' * (Limits().max_text_chars // 2 + 1) + '</p>'
            + byline(AUTHORS[10:], "dom") + FILE + '</article>')
    initial = document([], body=body)
    result = evaluate_page(initial, [expected], expected.id)
    assert result["extracted"]["partial"] is True
    assert result["status"] == (1 if complete else 2)


@pytest.fixture(scope="module")
def chromium():
    pytest.importorskip("playwright.sync_api", reason="Optional Playwright unavailable")
    with BrowserProbe(timeout_ms=7000, fixture_routes={}) as probe:
        if not probe.start():
            pytest.skip(probe.failure)
        yield probe


@pytest.mark.parametrize(("observed", "expected", "form", "status"), [
    (AUTHORS[:10], AUTHORS[:10], "dom", 1),
    (AUTHORS, AUTHORS[:10], "dom", 2),
    (AUTHORS, AUTHORS, "dom", 1),
    (list(reversed(AUTHORS)), AUTHORS, "dom", 1),
    (AUTHORS, AUTHORS, "ids", 1),
    (AUTHORS, AUTHORS[:10], "ids", 2),
], ids=["ten-equal", "eleven-v-ten", "eleven-equal", "reordered", "graph-ids", "graph-conflict"])
def test_real_chromium_compares_complete_sets(chromium, observed, expected, form, status):
    initial = document(observed, form)
    chromium.fixture_routes = {URL: (200, initial.body.decode("utf-8"))}
    rendered = chromium.check(URL)
    assert rendered.error_code == ""
    target = book(expected)
    assert evaluate_page(rendered, [target], target.id)["status"] == status


def test_cli_jsonl_xlsx_keep_complete_authors_and_row_decisions(tmp_path, monkeypatch):
    cases = [(AUTHORS[:10], AUTHORS[:10], "dom", 1), (AUTHORS, AUTHORS[:10], "dom", 2),
             (AUTHORS, AUTHORS, "dom", 1), (list(reversed(AUTHORS)), AUTHORS, "dom", 1),
             (AUTHORS, AUTHORS, "ids", 1), (AUTHORS, AUTHORS[:10], "ids", 2)]
    source = tmp_path / "author_input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet("Targets")
    sheet.append(HEADERS)
    responses = {}
    for index, (observed, expected, form, _) in enumerate(cases):
        url = URL + f'/{index}'
        sheet.append(("Publisher", "; ".join(expected), TITLE, "ebook", url, "offline.authors"))
        responses[url] = document(observed, form).body
    workbook.save(source)
    workbook.close()
    catalog, targets = load_inputs(source)
    monkeypatch.setattr(requests.sessions.Session, "get", lambda session, url, **kwargs: Session(responses[url]).get(url, **kwargs))
    output = tmp_path / "export"
    monkeypatch.setattr(sys, "argv", ["part2.py", "--input", str(source), "--browser", "off", "--delay", "0", "--output-dir", str(output)])
    main()
    records = [json.loads(line) for line in (output / "part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    workbook = openpyxl.load_workbook(output / "part2_results.xlsx", read_only=True)
    try:
        headers, *rows = list(workbook.active.values)
        cells = [dict(zip(headers, row)) for row in rows]
    finally:
        workbook.close()
    assert [record["status"] for record in records] == [case[3] for case in cases]
    assert len(records) == len(cells) == len(cases) and len(catalog) == 2
    for record, cell, target, (observed, _, _, _) in zip(records, cells, targets, cases):
        assert record["expected_book_id"] == target["expected_book_id"]
        assert authors_match(["; ".join(observed)], record["extracted"]["authors"])
        for key in ["status", "reason_code", "source_row", "url", "expected_book_id", "method"]:
            assert record[key] == cell[key]
        assert record["extracted"]["authors"] == json.loads(cell["authors"])
        assert record["content_signals"] == json.loads(cell["content_signals"])
