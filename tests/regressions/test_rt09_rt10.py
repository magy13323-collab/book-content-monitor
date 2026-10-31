"""Only RT-09/10: synthetic transport, unchanged audit inputs and isolated exports."""
from io import BytesIO
import json
from pathlib import Path
import sys

import openpyxl
import pytest
import requests

from browser import BrowserProbe
from pages import FetchResult, extract_page, fetch_url
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "rt09"
URL = 'https://audit.invalid/book'
BOOK = Book('target', 'Alice North', 'Silent Harbor')
HTML = '<article><h1>Silent Harbor</h1><p class="book-author">Alice North</p><p>Publisher: Harbor Press. Edition 2026. English language.</p>{}</article>'
FILE = '<a href="/files/123.epub">Download EPUB</a>'
ORIGINAL = (FIXTURES / 'json_conflict_in_removed_review.html').read_bytes()

def script(author='Bob West', **extra):
    return '<script type="application/ld+json">' + json.dumps({
        '@type': 'Book', 'name': BOOK.title, 'author': {'@type': 'Person', 'name': author}, **extra}) + '</script>'

def local(body, status=200, headers=None):
    return FetchResult(URL, URL, status, headers or {'Content-Type': 'text/html'},
                       body.encode() if isinstance(body, str) else body)

class Session:
    def __init__(self, status=200, body=b'', headers=None, error=None):
        self.status, self.body, self.error = status, body, error
        self.headers = {'Content-Type': 'text/html', **(headers or {})}
        self.calls = []

    def get(self, url, **kw):
        self.calls.append((url, kw))
        if self.error:
            raise self.error
        result = requests.Response()
        result.url, result.status_code = url, self.status
        result.headers.update(self.headers)
        result.raw = BytesIO(self.body)
        return result

    def close(self):
        pass

def evaluate(doc):
    return evaluate_page(doc, [BOOK], BOOK.id)

@pytest.mark.parametrize('case', ['json_conflict_in_removed_review', 'json_conflict_in_removed_review_scoped', 'own_file_positive'])
def test_original_rt09_oracle_is_unchanged(case):
    data = FIXTURES
    oracle = next(row['expected_statuses'] for row in json.loads((data / 'cases.json').read_text(encoding='utf-8')) if row['id'] == case)
    assert oracle == [1]
    assert evaluate(local((data / (case + '.html')).read_bytes()))['status'] in oracle

@pytest.mark.parametrize('region', ['aside', 'section', 'div', 'article'])
@pytest.mark.parametrize('reverse', [False, True])
def test_independent_review_metadata_is_retained_without_primary_ownership(region, reverse):
    own = HTML.format(FILE)
    review = f'<{region} class="reviews">{script()}</{region}>'
    doc = local(review + own if reverse else own + review)
    result = evaluate(doc)
    assert (result['status'], result['reason_code']) == (1, 'placement_signs')
    book = next(s for s in extract_page(doc).subjects if s['source'] == 'jsonld_book')
    assert book['authors'] == ['Bob West']
    assert not book['primary_metadata'] and not book['signals']

@pytest.mark.parametrize('body', [
    HTML.format(FILE + script()),
    HTML.format(FILE + '<aside class="reviews">' + script() + '</aside>'),
    HTML.format(FILE) + script(),
    HTML.format(FILE) + '<aside>' + script() + '</aside>',
    HTML.format(FILE) + '<aside class="reviews">' + script(mainEntityOfPage=URL) + '</aside>',
    HTML.format(FILE) + '<aside class="reviews">' + script(mainEntityOfPage={'@id': URL}) + '</aside>',
    HTML.format(FILE) + '<aside class="reviews">' + script(url=URL) + '</aside>',
    HTML.format(FILE) + '<aside class="reviews">' + script(**{'@id': URL}) + '</aside>',
    HTML.format(FILE) + '<aside class="reviews"><script type="application/ld+json">' + json.dumps({
        '@type': 'WebPage', '@id': URL, 'mainEntity': {'@type': 'Book', 'name': BOOK.title, 'author': {'name': 'Bob West'}}}) + '</script></aside>',
    '<script type="application/ld+json">' + json.dumps({'@type': 'WebPage', '@id': URL, 'mainEntity': {'@id': '#book'}}) + '</script>' + HTML.format(FILE) + '<aside class="reviews">' + script(**{'@id': '#book'}) + '</aside>',
    HTML.format(FILE + script('Alice North', **{'@id': '#book'})) + '<aside class="reviews">' + script(**{'@id': '#book'}) + '</aside>',
], ids=['own-card', 'nested-in-own-card', 'unknown-root', 'unknown-aside', 'page-string', 'page-reference', 'page-url', 'page-id', 'embedded-main-entity', 'inbound-main-entity', 'shared-book-id'])
def test_true_or_unresolved_ownership_still_protects_against_conflicts(body):
    result = evaluate(local(body))
    assert (result['status'], result['reason_code']) == (2, 'metadata_author_conflict')

def test_secondary_jsonld_cannot_supply_a_missing_primary_author_or_file():
    own = HTML.format('').replace('<p class="book-author">Alice North</p>', '')
    result = evaluate(local(own + '<aside class="reviews">' + script('Alice North') + FILE + '</aside>'))
    assert result['status'] != 1
    assert not result['content_signals']

@pytest.mark.parametrize('page_url', [None, 'https://other.invalid/review'])
def test_main_entity_of_an_independent_review_page_does_not_claim_the_target_page(page_url):
    page = {'@type': 'WebPage', 'mainEntity': {'@id': '#reviewed-book'}}
    if page_url:
        page['@id'] = page_url
    graph = {'@graph': [page, {'@id': '#reviewed-book', '@type': 'Book', 'name': BOOK.title, 'author': {'name': 'Bob West'}}]}
    body = HTML.format(FILE) + '<aside class="reviews"><script type="application/ld+json">' + json.dumps(graph) + '</script></aside>'
    result = evaluate(local(body))
    assert (result['status'], result['reason_code']) == (1, 'placement_signs')

@pytest.mark.parametrize('content', ['', FILE])
@pytest.mark.parametrize('range_value', [
    'bytes 0-{end}/9999', 'bytes 10-{end}/9999', 'bytes 0-{end}/*',
    None, '', 'broken', 'bytes */9999', 'bytes 20-10/9999',
    'bytes 0-9999/10', 'bytes 0-{end}/{end}', 'bytes 0-{size}/{size}',
], ids=['partial', 'nonzero-start', 'unknown-total', 'absent', 'empty', 'malformed', 'unsatisfied', 'reversed', 'end-over-total', 'off-by-one', 'body-mismatch'])
def test_rt10_incomplete_transport_is_not_a_negative_but_keeps_positive(content, range_value):
    body = HTML.format(content).encode()
    headers = {} if range_value is None else {'Content-Range': range_value.format(end=len(body)-1, size=len(body))}
    session = Session(206, body, headers)
    doc = fetch_url(URL, session=session)
    assert doc.http_status == 206 and not doc.error_code
    assert doc.headers.get('content-range') == headers.get('Content-Range')
    assert len(session.calls) == 1 and 'Range' not in session.calls[0][1]['headers']
    result = evaluate(doc)
    assert result['extracted']['partial'] is True
    assert (result['status'], result['reason_code']) == ((1, 'placement_signs') if content else (2, 'partial_document'))

@pytest.mark.parametrize('status', [200, 206])
def test_complete_range_and_normal_200_remain_negative(status):
    body = HTML.format('').encode()
    doc = fetch_url(URL, session=Session(status, body, {'Content-Range': f'bytes 0-{len(body)-1}/{len(body)}'} if status == 206 else {}))
    result = evaluate(doc)
    assert (result['status'], result['reason_code'], result['extracted']['partial']) == (0, 'content_not_confirmed', False)

def test_direct_fetch_result_with_mixed_case_range_headers_has_the_same_partial_semantics():
    result = evaluate(local(HTML.format('').encode(), 206, {'cOnTeNt-TyPe': 'text/html', 'cOnTeNt-RaNgE': 'bytes 0-141/9999'}))
    assert result['status'] == 2 and result['extracted']['partial']

def test_partial_conflicting_target_does_not_become_positive():
    body = HTML.format(FILE + script()).encode()
    result = evaluate(fetch_url(URL, session=Session(206, body, {'Content-Range': f'bytes 0-{len(body)-1}/9999'})))
    assert result['extracted']['partial']
    assert (result['status'], result['reason_code']) == (2, 'metadata_author_conflict')

def test_partial_fragment_identifying_only_another_book_cannot_prove_absence():
    body = HTML.format(FILE).replace('Silent Harbor', 'Winter Garden').replace('Alice North', 'Bob West').encode()
    result = evaluate(fetch_url(URL, session=Session(206, body)))
    assert (result['status'], result['reason_code']) == (2, 'partial_document')

@pytest.mark.parametrize('status,expected', [(200, 2), (204, 2), (206, 2), (404, 0), (410, 0), (403, 2), (429, 2), (500, 2), (503, 2)])
def test_empty_and_http_error_controls(status, expected):
    result = evaluate(fetch_url(URL, session=Session(status)))
    assert result['status'] == expected
    assert result['extracted']['partial'] is (status == 206)

@pytest.mark.parametrize('error,code', [(requests.Timeout('offline'), 'timeout'), (requests.exceptions.SSLError('offline'), 'tls_error'), (requests.ConnectionError('offline'), 'network_error')])
def test_load_errors_remain_uncertain(error, code):
    result = evaluate(fetch_url(URL, session=Session(error=error)))
    assert (result['status'], result['reason_code']) == (2, code)

@pytest.fixture(scope='module')
def chromium():
    with BrowserProbe(timeout_ms=7000, fixture_routes={}) as probe:
        assert probe.start(), probe.failure
        yield probe

@pytest.mark.parametrize('body,status', [(ORIGINAL, 1), (HTML.format(FILE + script()).encode(), 2)])
def test_real_chromium_preserves_rt09_ownership(chromium, body, status):
    chromium.fixture_routes = {URL: (200, body.decode())}
    doc = chromium.check(URL)
    assert not doc.error_code and doc.visibility_checked
    assert evaluate(doc)['status'] == status

def test_public_cli_jsonl_xlsx_preserves_rt09_rt10_and_transport_controls(tmp_path, monkeypatch):
    cases = [(200, ORIGINAL, {}, 1),
             (200, HTML.format(FILE + script()).encode(), {}, 2),
             (206, HTML.format('').encode(), {'Content-Range': 'bytes 0-141/9999'}, 2),
             (200, HTML.format('').encode(), {}, 0),
             (206, HTML.format(FILE).encode(), {}, 1),
             (206, b'', {}, 2), (204, b'', {}, 2), (404, b'', {}, 0), (410, b'', {}, 0),
             (503, b'', {}, 2), (200, HTML.format(FILE).encode(), {}, 1)]
    source = tmp_path / 'input.xlsx'
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet('Targets')
    sheet.append(HEADERS)
    routes = {}
    for index, (status, body, headers, _) in enumerate(cases):
        url = URL + f'/{index}'
        sheet.append(['Publisher', BOOK.author, BOOK.title, 'ebook', url, 'audit.invalid'])
        routes[url] = Session(status, body, headers)
    workbook.save(source)
    workbook.close()
    catalog, targets = load_inputs(source)
    monkeypatch.setattr(requests.sessions.Session, 'get', lambda self, url, **kw: routes[url].get(url, **kw))
    dest = tmp_path / 'cli_export'
    monkeypatch.setattr(sys, 'argv', ['part2.py', '--input', str(source), '--browser', 'off', '--delay', '0', '--output-dir', str(dest)])
    main()
    records = [json.loads(line) for line in (dest / 'part2_results.jsonl').read_text(encoding='utf-8').splitlines()]
    workbook = openpyxl.load_workbook(dest / 'part2_results.xlsx', read_only=True)
    try:
        headers, *rows = list(workbook.active.values)
        cells = [dict(zip(headers, row)) for row in rows]
    finally:
        workbook.close()
    assert [row['status'] for row in records] == [case[-1] for case in cases]
    assert len(records) == len(cells) == len(targets) and len(catalog) == 1
    for row, cell, target in zip(records, cells, targets):
        assert row['expected_book_id'] == target['expected_book_id'] == catalog[0].id
        for key in ['status', 'reason_code', 'url', 'source_row', 'expected_book_id', 'method', 'checked_at', 'partial']:
            assert (row['extracted']['partial'] if key == 'partial' else row[key]) == cell[key]
        for key in ['evidence', 'content_signals', 'observed_formats']:
            assert row[key] == json.loads(cell[key])
