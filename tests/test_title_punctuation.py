"""Punctuation is evidence; equal words alone cannot establish title identity."""

import hashlib
import json

import openpyxl
import pytest
import requests

from pages import FetchResult, extract_page
from part1 import process_workbook, write_workbook
from part2 import Book, analyze_url, evaluate_page, load_inputs, write_results
from test_part2 import synthetic_pdf
from test_yo_identity import document, source


A = "Казнить нельзя, помиловать"
B = "Казнить, нельзя помиловать"
AUTHOR = "Анна Лесная"
BOOKS = [Book("a", AUTHOR, A), Book("b", AUTHOR, B)]


@pytest.fixture(autouse=True)
def no_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def test_catalogue_ids_already_preserve_commas_and_row_order_does_not_matter(tmp_path):
    rows = [(AUTHOR, A), (AUTHOR, B)]
    path = source(tmp_path, rows)
    analysis = process_workbook(path)
    catalog, targets = load_inputs(path)
    expected = {A: "book_a949cd5d76d93dfe", B: "book_f9d30b23004511a9"}
    assert len(analysis.records) == len(analysis.books) == len(catalog) == 2
    assert {book.title: book.id for book in catalog} == expected
    assert [record.book_id for record in analysis.records] == [target["expected_book_id"] for target in targets]
    reversed_catalog, _ = load_inputs(source(tmp_path, list(reversed(rows)), "reversed.xlsx"))
    assert {book.title: book.id for book in reversed_catalog} == expected


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize("observed", [A, B])
@pytest.mark.parametrize("book", BOOKS, ids=["a", "b"])
def test_semantic_comma_identity_matrix(book, observed, structured):
    result = evaluate_page(document(observed, structured=observed if structured else None), [book], book.id)
    assert (result["status"], result["reason_code"]) == (
        (1, "placement_signs") if book.title == observed else (2, "ambiguous"))
    if book.title != observed:
        assert not result["content_signals"]
        assert any("пунктуац" in item.lower() for item in result["evidence"])


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize("book", BOOKS, ids=["a", "b"])
def test_same_author_and_words_do_not_confuse_catalogue_selection(book, structured):
    result = evaluate_page(document(book.title, structured=book.title if structured else None), BOOKS)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")
    assert result["match"]["book_id"] == book.id
    assert not result["match"]["other_catalog_matches"]


@pytest.mark.parametrize(("expected", "observed"), [
    ("Казнить: нельзя помиловать", "Казнить нельзя: помиловать"),
    ("Казнить; нельзя помиловать", "Казнить нельзя; помиловать"),
    ("Казнить? нельзя помиловать", "Казнить нельзя? помиловать"),
    ("Казнить! нельзя помиловать", "Казнить нельзя! помиловать"),
    ("Казнить. Нельзя помиловать", "Казнить нельзя. Помиловать"),
    ("Казнить — нельзя помиловать", "Казнить нельзя — помиловать"),
    ("(Казнить нельзя) помиловать", "Казнить (нельзя помиловать)"),
])
def test_other_internal_punctuation_differences_are_uncertain(expected, observed):
    result = evaluate_page(document(observed), [Book("target", AUTHOR, expected)], "target")
    assert (result["status"], result["reason_code"]) == (2, "ambiguous")
    assert not result["content_signals"]


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize(("expected", "observed"), [
    (A, "  КАЗНИТЬ\u00a0 нельзя ,   помиловать  "),
    (A, '«Казнить нельзя, помиловать»'),
    ('«Казнить» нельзя, помиловать', '"Казнить" нельзя, помиловать'),
    ('Казнить — нельзя помиловать', 'Казнить - нельзя помиловать'),
    ('Казнить–нельзя помиловать', 'Казнить — нельзя помиловать'),
    (A, 'Читать «Казнить нельзя, помиловать» онлайн'),
])
def test_allowed_presentation_differences_keep_positive(expected, observed, structured):
    book = Book("target", AUTHOR, expected)
    result = evaluate_page(document(observed, structured=expected if structured else None), [book], book.id)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")


@pytest.mark.parametrize(("dom", "metadata"), [(A, B), (B, A)])
def test_jsonld_does_not_inherit_a_different_comma_titles_file(dom, metadata):
    initial = document(dom, structured=metadata)
    result = evaluate_page(initial, [Book("target", AUTHOR, metadata)], "target")
    assert (result["status"], result["reason_code"]) == (2, "metadata_title_conflict")
    assert not result["content_signals"]
    assert not next(s for s in extract_page(initial).subjects if s["source"] == "jsonld_book")["signals"]


@pytest.mark.parametrize(("title", "label"), [(A, B), (B, A)])
def test_descriptive_link_punctuation_does_not_prove_ownership(title, label):
    initial = document(title, content=f'<a href="/files/123.epub">Скачать {label} EPUB</a>')
    result = evaluate_page(initial, [Book("target", AUTHOR, title)], "target")
    assert (result["status"], result["reason_code"]) == (2, "content_scope_unresolved")
    assert not result["content_signals"]


@pytest.mark.parametrize(("primary", "nested"), [(A, B), (B, A)])
def test_nested_title_with_different_commas_cannot_supply_an_own_file(primary, nested):
    initial = document(primary, content=f'<article><h2>{nested}</h2><p class="book-author">{AUTHOR}</p>'
                       '<a href="/files/123.epub">Скачать EPUB</a></article>')
    result = evaluate_page(initial, [Book("target", AUTHOR, primary)], "target")
    assert (result["status"], result["reason_code"]) == (2, "ambiguous")
    assert not result["content_signals"]
    assert any(s["binding"] == "unresolved" for s in result["extracted"]["observed_content_signals"])


def test_exact_own_file_survives_next_to_a_punctuation_variant():
    initial = document(A, content='<a href="/files/own.epub">Скачать EPUB</a>'
                       f'<article><h2>{B}</h2><p class="book-author">{AUTHOR}</p>'
                       '<a href="/files/other.epub">Скачать EPUB</a></article>')
    result = evaluate_page(initial, [BOOKS[0]], "a")
    assert result["status"] == 1
    assert [s["url"] for s in result["content_signals"]] == ["https://offline.test/files/own.epub"]


@pytest.mark.parametrize(("opening", "metadata", "status", "reason"), [
    ("Execute not, pardon", "Execute not, pardon", 1, "placement_signs"),
    ("Execute, not pardon", "Execute, not pardon", 2, "ambiguous"),
    ("Execute, not pardon", "Execute not, pardon", 2, "pdf_identity_unresolved"),
])
def test_pdf_identity_uses_the_same_punctuation_evidence(opening, metadata, status, reason):
    initial = FetchResult("https://offline.test/book.pdf", "https://offline.test/book.pdf", 200,
                          {"Content-Type": "application/pdf"}, synthetic_pdf(
                              title=opening, metadata={"/Title": metadata, "/Author": "Alice Example"}))
    result = evaluate_page(initial, [Book("target", "Alice Example", "Execute not, pardon")], "target")
    assert (result["status"], result["reason_code"]) == (status, reason)


@pytest.mark.parametrize("crossed", [False, True], ids=["exact", "uncertain"])
def test_exports_preserve_source_rows_ids_and_decisions(tmp_path, crossed):
    path = source(tmp_path, [(AUTHOR, A), (AUTHOR, B)])
    original = hashlib.sha256(path.read_bytes()).digest()
    analysis = process_workbook(path)
    catalog, targets = load_inputs(path)
    write_workbook(analysis, tmp_path / "processed.xlsx")
    results = []
    for target in targets:
        title = target["expected_title"]
        observed = (B if title == A else A) if crossed else title
        initial = document(observed)
        results.append(analyze_url(target, catalog, fetcher=lambda url, initial=initial: initial))
    write_results(results, tmp_path / "output")
    exported = [json.loads(line) for line in (tmp_path / "output/part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["source_row"] for r in exported] == [2, 3]
    assert [r["expected_title"] for r in exported] == [A, B]
    assert [r["expected_book_id"] for r in exported] == ["book_a949cd5d76d93dfe", "book_f9d30b23004511a9"]
    assert [r["status"] for r in exported] == ([2, 2] if crossed else [1, 1])
    workbook = openpyxl.load_workbook(tmp_path / "processed.xlsx", read_only=True)
    assert workbook["Normalized"].max_row == workbook["Books"].max_row == 3
    assert [(r[2], r[6]) for r in list(workbook["Normalized"].values)[1:]] == [(A, 2), (B, 3)]
    workbook.close()
    assert hashlib.sha256(path.read_bytes()).digest() == original
