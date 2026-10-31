"""Independent RT-01 controls: URL fragments are not file destinations."""

import hashlib
from io import BytesIO
import json
import sys
from urllib.parse import urljoin

import openpyxl
import pytest
import requests

from browser import BrowserProbe
from pages import FetchResult
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main


BOOK = Book("target", "Alice North", "Silent Harbor")
OTHER = Book("other", "Bob West", "Winter Garden")
URL = "https://offline.fragments/book"
DETAILS = "<p>Publication details: language English; publisher Harbor Press.</p>"
ANCHOR = '<a href="#book.epub">Download EPUB</a>'
FILE = '<a href="/files/book.epub">Download EPUB</a>'
READER = '<div class="reader-text">' + (
    "The traveler followed the path beside the harbor and reached the old house before sunset. " * 18
) + "</div>"


@pytest.fixture(autouse=True)
def no_live_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden in RT-01 regressions")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def card(content, book=BOOK, heading="h1", author=True):
    byline = f'<p class="book-author">{book.author}</p>' if author else ""
    return f"<article><{heading}>{book.title}</{heading}>{byline}{DETAILS}{content}</article>"


def document(content, structured=False, url=URL, foreign=False):
    metadata = ""
    if structured:
        metadata = '<script type="application/ld+json">' + json.dumps({
            "@context": "https://schema.org", "@type": "Book", "name": BOOK.title,
            "author": {"@type": "Person", "name": BOOK.author},
        }) + "</script>"
    body = card("" if foreign else content, author=not structured)
    if foreign:
        body += card(content, book=OTHER, heading="h2")
    html = f'<html><head><meta charset="utf-8">{metadata}</head><body><main>{body}</main></body></html>'
    return FetchResult(url, url, 200, {"Content-Type": "text/html; charset=utf-8"}, html.encode())


@pytest.mark.parametrize("structured", [False, True], ids=["dom", "jsonld"])
@pytest.mark.parametrize(("href", "status"), [
    ("#book.epub", 0),
    ("#book.pdf", 0),
    ("/viewer#book.pdf", 0),
    ("/files/book.epub", 1),
    ("/files/book.pdf", 1),
    ("/files/book.epub#chapter", 1),
    ("/files/book.pdf#page=2", 1),
    ("/files/book.epub?download=1", 1),
    ("/files/book.pdf?download=1#page=2", 1),
])
def test_fragment_is_not_a_file_but_real_path_still_is(href, status, structured):
    extension = "PDF" if ".pdf" in href else "EPUB"
    result = evaluate_page(document(f'<a href="{href}">Download {extension}</a>', structured), [BOOK], BOOK.id)
    assert result["status"] == status
    assert result["match"]["state"] == "matched"
    if status == 1:
        assert result["reason_code"] == "placement_signs"
        assert [signal["url"] for signal in result["content_signals"]] == [urljoin(URL, href)]
        assert result["content_signals"][0]["binding"] == "bound"
        assert result["content_signals"][0]["scope_title"] == BOOK.title
    else:
        assert result["reason_code"] == "content_not_confirmed"
        assert result["content_signals"] == []
        assert result["extracted"]["observed_content_signals"] == []


@pytest.mark.parametrize("structured", [False, True], ids=["dom", "jsonld"])
def test_foreign_epub_remains_foreign(structured):
    initial = document(FILE, structured, foreign=True)
    own = evaluate_page(initial, [BOOK, OTHER], BOOK.id)
    other = evaluate_page(initial, [BOOK, OTHER], OTHER.id)
    assert (own["status"], own["content_signals"]) == (0, [])
    assert other["status"] == 1
    assert [signal["scope_title"] for signal in other["content_signals"]] == [OTHER.title]


@pytest.mark.parametrize("structured", [False, True], ids=["dom", "jsonld"])
@pytest.mark.parametrize(("positive", "kind"), [(FILE, "file_link"), (READER, "reader_text")], ids=["own-file", "own-text"])
def test_anchor_does_not_cancel_independent_positive_evidence(structured, positive, kind):
    result = evaluate_page(document(ANCHOR + positive, structured), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")
    assert [signal["kind"] for signal in result["content_signals"]] == [kind]
    assert all(signal["raw_resource"] != "#book.epub" for signal in result["content_signals"])


@pytest.mark.parametrize("fragment", ["", "#chapter"])
def test_rt12_query_policy_is_unchanged(fragment):
    # Compatibility guard only: RT-12 policy has not been approved for a change.
    href = "/users/view?return=book.epub" + fragment
    result = evaluate_page(document(f'<a href="{href}">Download EPUB</a>'), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert result["content_signals"][0]["url"] == urljoin(URL, href)


@pytest.fixture(scope="module")
def chromium():
    pytest.importorskip("playwright.sync_api", reason="Optional Playwright unavailable")
    with BrowserProbe(timeout_ms=7000, fixture_routes={}) as probe:
        if not probe.start():
            pytest.skip(probe.failure)
        yield probe


@pytest.mark.parametrize(("content", "foreign", "status"), [
    (ANCHOR, False, 0),
    (FILE, False, 1),
    ('<a href="/files/book.epub#chapter">Download EPUB</a>', False, 1),
    ('<a href="/files/book.epub?download=1">Download EPUB</a>', False, 1),
    (FILE, True, 0),
    (ANCHOR + READER, False, 1),
], ids=["fragment-only", "own-file", "file-fragment", "file-query", "foreign-file", "anchor-with-reader"])
def test_real_chromium_keeps_rt01_decision_and_ownership(chromium, content, foreign, status):
    initial = document(content, structured=True, foreign=foreign)
    chromium.fixture_routes = {URL: (200, initial.body.decode("utf-8"))}
    rendered = chromium.check(URL)
    assert rendered.error_code == ""
    assert rendered.visibility_checked is True
    result = evaluate_page(rendered, [BOOK], BOOK.id)
    assert result["status"] == status
    assert all(signal["raw_resource"] != "#book.epub" for signal in result["content_signals"])


def test_public_xlsx_cli_http_and_both_exports_keep_rt01_results(tmp_path, monkeypatch):
    scenarios = [
        (ANCHOR, False, 0),
        (FILE, False, 1),
        ('<a href="/files/book.epub#chapter">Download EPUB</a>', False, 1),
        ('<a href="/files/book.epub?download=1">Download EPUB</a>', False, 1),
        (FILE, True, 0),
        (ANCHOR + READER, False, 1),
    ]
    source = tmp_path / "rt01_input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet("Targets")
    sheet.append(HEADERS)
    responses = {}
    for index, (content, foreign, _) in enumerate(scenarios):
        url = URL + f"/{index}"
        sheet.append(("Publisher", BOOK.author, BOOK.title, "ebook", url, "offline.fragments"))
        responses[url] = document(content, structured=True, url=url, foreign=foreign).body
    workbook.save(source)
    workbook.close()
    original_hash = hashlib.sha256(source.read_bytes()).digest()
    catalog, targets = load_inputs(source)
    calls = []

    def synthetic_get(session, url, **kwargs):
        calls.append(url)
        assert kwargs["allow_redirects"] is False
        response = requests.Response()
        response.status_code = 200
        response.url = url
        response.headers["Content-Type"] = "text/html; charset=utf-8"
        response.headers["Content-Length"] = str(len(responses[url]))
        response.raw = BytesIO(responses[url])
        return response

    monkeypatch.setattr(requests.sessions.Session, "get", synthetic_get)
    output = tmp_path / "export"
    monkeypatch.setattr(sys, "argv", ["part2.py", "--input", str(source), "--browser", "off", "--delay", "0", "--output-dir", str(output)])
    main()
    records = [json.loads(line) for line in (output / "part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    workbook = openpyxl.load_workbook(output / "part2_results.xlsx", read_only=True)
    try:
        headers, *rows = list(workbook.active.values)
        cells = [dict(zip(headers, row)) for row in rows]
    finally:
        workbook.close()
    assert calls == list(responses)
    assert len(records) == len(cells) == len(scenarios)
    assert [record["status"] for record in records] == [case[2] for case in scenarios]
    assert len(catalog) == 1
    for record, cell, target in zip(records, cells, targets):
        assert record["expected_book_id"] == target["expected_book_id"] == catalog[0].id
        for field in ["source_row", "url", "expected_book_id", "status", "reason_code", "method", "checked_at"]:
            assert record[field] == cell[field]
        assert record["source_row"] == target["source_row"]
        for field in ["content_signals", "evidence", "observed_formats"]:
            assert record[field] == json.loads(cell[field])
        assert all(signal["raw_resource"] != "#book.epub" for signal in record["content_signals"])
    assert hashlib.sha256(source.read_bytes()).digest() == original_hash
