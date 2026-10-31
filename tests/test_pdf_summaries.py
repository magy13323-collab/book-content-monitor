"""RT-08: explicit PDF summaries are secondary; narrative mentions remain valid."""

import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys

import openpyxl
from pypdf import PdfReader, PdfWriter
import pytest
import requests

from pages import FetchResult, _pdf_document_role
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main
from test_part2 import synthetic_pdf
from test_pdf_roles import AUTHOR, BOOK, TITLE, assert_secondary, document, evaluate
from test_partial_pdf import RangeResponse, local, network, pdf as partial_pdf
from test_short_pdf import pdf as short_pdf


FIXTURE = Path(__file__).parent / "fixtures/rt08/pdf_summary_role.pdf"
SHA256 = "fc0b3378a709a76828a04f08468639de93ee1e1952491ea5c5cb045fbe59a312"
ORIGINAL_BOOK = Book("target", "Alice North", "Silent Harbor")
SUMMARY_BODY = "This summary describes the plot of the novel and discusses the choices of its main characters."


@pytest.fixture(autouse=True)
def no_live_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden in RT-08 regressions")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def original():
    body = FIXTURE.read_bytes()
    assert hashlib.sha256(body).hexdigest() == SHA256
    return FetchResult("https://audit.invalid/book", "https://audit.invalid/book", 200,
                       {"Content-Type": "application/pdf"}, body)


def metadata(subject):
    return {"/Title": TITLE, "/Author": AUTHOR, "/Subject": subject}


def test_original_audit_pdf_keeps_unchanged_oracle():
    result = evaluate_page(original(), [ORIGINAL_BOOK], ORIGINAL_BOOK.id)
    assert result["status"] in {0, 2}  # Original independent oracle is not changed to require 1.
    assert (result["status"], result["reason_code"]) == (2, "pdf_document_role_unresolved")
    assert not result["content_signals"]
    assert result["extracted"]["pdf_pages_read"] == result["extracted"]["pdf_page_count"] == 2
    assert result["extracted"]["metadata"]["document_role"] == "summary"
    assert result["extracted"]["metadata"]["document_role_evidence"] == "Summary of the novel"
    assert any("Summary of the novel" in line for line in result["evidence"])
    assert "пересказа" in result["reason"]


@pytest.mark.parametrize("label", [
    "Summary of the novel", "A summary of the book.", "Summary of the work: The Northern Wind",
    "Plot summary", "Plot-summary of the novel", "This is a summary of the novel.",
])
@pytest.mark.parametrize("location", ["subject", "opening", "both"])
def test_explicit_summary_roles_in_real_pdf(label, location):
    pdf = document([label] if location != "subject" else ["Chapter 1"],
                   metadata=metadata(label if location != "opening" else ""), body=SUMMARY_BODY)
    assert_secondary(evaluate(pdf), "summary")


@pytest.mark.parametrize("label", ["Краткое содержание романа", "Пересказ книги", "Краткое изложение произведения"])
def test_unicode_subject_identifies_a_summary(label):
    assert_secondary(evaluate(document(["Chapter 1"], metadata=metadata(label))), "summary")


def test_long_summary_does_not_become_a_book_through_page_count():
    source = PdfReader(BytesIO(document(["Summary of the novel"], body=SUMMARY_BODY).body))
    writer = PdfWriter()
    writer.add_metadata(metadata("Summary of the novel"))
    for index in range(10):
        writer.add_page(source.pages[index % 2])
    output = BytesIO()
    writer.write(output)
    result = evaluate(FetchResult("https://offline.test/book.pdf", "https://offline.test/book.pdf", 200,
                                 {"Content-Type": "application/pdf"}, output.getvalue()))
    assert (result["status"], result["reason_code"]) == (2, "pdf_document_role_unresolved")
    assert result["extracted"]["pdf_pages_read"] == result["extracted"]["pdf_page_count"] == 10
    assert result["extracted"]["text_chars"] > 20000 and result["extracted"]["partial"] is False
    assert not result["content_signals"]


@pytest.mark.parametrize("mention", [
    "Summary", "Summary of the novel", "Plot summary", "Summary of the novel was on her desk.",
    "She read a summary of the novel and continued her journey.",
])
def test_mentions_in_a_chapter_do_not_reject_the_book(mention):
    result = evaluate(document(["Chapter 1", mention]))
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")
    assert "document_role" not in result["extracted"]["metadata"]


@pytest.mark.parametrize("subject", [
    "Summary", "Summary statistics for the harbor", "A novel about a summary",
    "The word summary occurs in a character's letter.", "Summary of the novel was on her desk.",
    "Summary of the novel, she thought, had told her nothing about the stranger.",
])
def test_subject_topic_or_narrative_is_not_a_document_role(subject):
    result = evaluate(document(["Chapter 1"], metadata=metadata(subject)))
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")
    assert "document_role" not in result["extracted"]["metadata"]


@pytest.mark.parametrize("mention", [
    "Summary", "Summary statistics were pinned beside the door.",
    "Summary of the novel was on her desk.", "He unfolded a summary of the novel and read it.",
    "Summary of the novel, she thought, had told her nothing about the stranger.",
])
def test_narrative_before_a_chapter_is_not_an_explicit_summary_label(mention):
    assert evaluate(document([mention]))["status"] == 1


@pytest.mark.parametrize("title", ["The Summary", "Summary of the Novel"])
def test_matching_book_title_is_identity_not_a_secondary_role(title):
    body = synthetic_pdf(title, AUTHOR, metadata={"/Title": title, "/Author": AUTHOR}, note="Chapter 1")
    book = Book("summary-title", AUTHOR, title)
    pdf = FetchResult("https://offline.test/book.pdf", "https://offline.test/book.pdf", 200,
                      {"Content-Type": "application/pdf"}, body)
    result = evaluate_page(pdf, [book], book.id)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")


def test_keywords_do_not_declare_a_pdf_subject_role():
    result = evaluate(document(["Chapter 1"], metadata={**metadata("A novel"), "/Keywords": "Summary of the novel"}))
    assert result["status"] == 1


@pytest.mark.parametrize("text", ["Summary", "Краткое содержание", "Обзор", "Summary of the novel", "Пересказ книги"])
def test_role_words_inside_unicode_narrative_keep_chapter_context(text):
    role, evidence = _pdf_document_role(f"{TITLE}\n{AUTHOR}\nChapter 1\n{text}\nThe traveler continued north.", "",
                                        {TITLE.casefold(), AUTHOR.casefold()})
    assert (role, evidence) == ("", "")


def test_document_level_summary_statement_after_a_chapter_stays_uncertain():
    assert_secondary(evaluate(document(["Chapter 1", "This document provides a summary of the novel."])), "summary")


def test_opening_role_outweighs_a_conflicting_subject_claim():
    assert_secondary(evaluate(document(["Summary of the novel"], metadata=metadata("Original novel"))), "summary")


def test_missing_metadata_does_not_restore_a_summary_signal():
    result = evaluate(document(["Summary of the novel"], metadata={}))
    assert (result["status"], result["reason_code"]) == (2, "pdf_document_role_unresolved")
    assert not result["content_signals"]


def test_identity_conflict_keeps_its_existing_priority():
    result = evaluate(document(["Summary of the novel"], metadata={
        "/Title": "South Road", "/Author": "Bob Different", "/Subject": "Summary of the novel"}))
    assert (result["status"], result["reason_code"]) == (2, "pdf_identity_unresolved")
    assert not result["content_signals"]
    assert result["extracted"]["metadata"]["document_role"] == "summary"


@pytest.mark.parametrize("mode", ["local", "range"])
@pytest.mark.parametrize("summary", [False, True])
def test_partial_pdf_preserves_book_or_secondary_role(mode, summary):
    body = partial_pdf(own=True, note="Summary of the novel" if summary else "Chapter 1", padding=mode == "range")
    fetched = local(body) if mode == "local" else network(body)[0]
    result = evaluate(fetched)
    assert result["extracted"]["partial"] is True
    assert result["extracted"]["pdf_pages_read"] == (20 if mode == "local" else 3)
    assert (result["status"], result["reason_code"]) == ((2, "pdf_document_role_unresolved") if summary else (1, "placement_signs"))
    assert bool(result["content_signals"]) is not summary


@pytest.mark.parametrize("first_chars,pages,status", [(666, 1, 2), (1544, 1, 2), (666, 2, 2), (1544, 2, 1)])
@pytest.mark.parametrize("summary", [False, True])
def test_short_pdf_thresholds_are_unchanged(first_chars, pages, status, summary):
    result = evaluate(short_pdf(first_chars, pages, label="Summary of the novel" if summary else "Chapter 1"))
    assert result["status"] == (2 if summary else status)
    assert result["reason_code"] == ("pdf_document_role_unresolved" if summary else
                                      "placement_signs" if status == 1 else "content_not_confirmed")
    assert result["extracted"]["partial"] is False


def test_cli_jsonl_xlsx_preserve_pdf_role_and_evidence(tmp_path, monkeypatch):
    cases = [
        (ORIGINAL_BOOK, original(), 2, "pdf_document_role_unresolved"),
        (BOOK, document(["Chapter 1"]), 1, "placement_signs"),
        (BOOK, document(["Chapter 1"], metadata=metadata("Summary of the novel")), 2, "pdf_document_role_unresolved"),
        (BOOK, document(["Summary of the novel"], body=SUMMARY_BODY), 2, "pdf_document_role_unresolved"),
        (BOOK, document(["Chapter 1", "Summary of the novel"]), 1, "placement_signs"),
        (BOOK, document(["Chapter 1"], metadata=metadata("Summary")), 1, "placement_signs"),
        (BOOK, document(metadata={"/Title": "South Road", "/Author": "Bob Different"}), 2, "pdf_identity_unresolved"),
        (BOOK, short_pdf(), 2, "content_not_confirmed"),
        (BOOK, local(partial_pdf(own=True, note="Summary of the novel")), 2, "pdf_document_role_unresolved"),
        (BOOK, local(partial_pdf(own=True, note="Chapter 1")), 1, "placement_signs"),
    ]
    source = tmp_path / "rt08.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet("Targets")
    sheet.append(HEADERS)
    responses = {}
    for i, (book, pdf, _, _) in enumerate(cases):
        url = f"https://offline.summaries/{i}.pdf"
        sheet.append(("Publisher", book.author, book.title, "ebook", url, "offline.summaries"))
        responses[url] = pdf.body
    workbook.save(source)
    workbook.close()
    source_sha = hashlib.sha256(source.read_bytes()).digest()
    _, targets = load_inputs(source)
    calls = []

    def synthetic_get(session, url, **kwargs):
        calls.append(url)
        requested = re.fullmatch(r"bytes=(\d+)-(\d+)", kwargs["headers"]["Range"])
        assert requested and kwargs["allow_redirects"] is False
        body = responses[url]
        start, end = int(requested[1]), min(int(requested[2]), len(body) - 1)
        data = body[start:end + 1]
        return RangeResponse(206, data, {"Content-Type": "application/pdf", "Content-Length": str(len(data)),
                                         "Content-Range": f"bytes {start}-{end}/{len(body)}"})

    monkeypatch.setattr(requests.sessions.Session, "get", synthetic_get)
    output = tmp_path / "export"
    monkeypatch.setattr(sys, "argv", ["part2.py", "--input", str(source), "--browser", "off",
                                     "--delay", "0", "--output-dir", str(output)])
    main()
    records = [json.loads(line) for line in (output / "part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    workbook = openpyxl.load_workbook(output / "part2_results.xlsx", read_only=True)
    headers, *rows = list(workbook.active.values)
    workbook.close()
    assert calls == list(responses)
    assert len(records) == len(rows) == len(cases)
    for record, row, target, (_, _, status, reason) in zip(records, rows, targets, cases):
        cells = dict(zip(headers, row))
        assert (record["status"], record["reason_code"]) == (status, reason)
        assert record["expected_book_id"] == target["expected_book_id"]
        for key in ["status", "reason_code", "expected_book_id", "source_row", "url", "method"]:
            assert record[key] == cells[key]
        assert record["evidence"] == json.loads(cells["evidence"])
        assert record["content_signals"] == json.loads(cells["content_signals"])
        assert record["extracted"]["metadata"] == json.loads(cells["metadata"])
        if reason == "pdf_document_role_unresolved":
            assert not record["content_signals"]
            marker = record["extracted"]["metadata"]["document_role_evidence"]
            assert marker and any(marker in item for item in record["evidence"])
    assert hashlib.sha256(source.read_bytes()).digest() == source_sha
