"""RT-05/06: scoped incomplete reading evidence is uncertainty, never proof."""

import hashlib
from html import escape
from io import BytesIO
import json
from pathlib import Path
import sys

import openpyxl
import pytest
import requests

from browser import BrowserProbe
from pages import FetchResult
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main


BOOK = Book("target", "Alice North", "Silent Harbor")
OTHER = Book("other", "Bob West", "Winter Garden")
URL = "https://offline.reader/book"
FIXTURES = Path(__file__).parent / "fixtures" / "rt05_rt06"
ORIGINALS = [
    ("reader_unusual_wrapper", "61d509d310000dd7e858ef329d049288185deffab5d3cd33ea7122cd18e342ac"),
    ("reader_semantic_section", "717dba219f919f7a9c91799553bdb124cd995cdee53804a35ca40789d3ce8680"),
    ("reader_short", "88bf5f225984bcb3fe7428bfd34427d8597dd36994b1e0b8915bfcdc1266b16b"),
    ("embedded_iframe_complete_card", "a9acec6cd204c2e1ea83dce7faa2f1bae7e79ae9838131eb5f06f21afe7d77ed"),
    ("embedded_object_complete_card", "686ead6ec8dadd353899623133ad19b449dd56db567239cab47b3a972fa1a950"),
    ("embedded_embed_complete_card", "c124296c557dd2cbad1531ea8ec50573166ae39d19d2546e82b0b69353007827"),
]
PROSE = "The lantern beside the harbor lit the narrow path as she returned to the empty house. " * 18
READER = f'<div class="reader-text"><h2>Chapter 1</h2><p>{PROSE}</p></div>'
SHORT = '<div class="reader-text"><p>An explicitly marked short excerpt.</p></div>'
SEMANTIC = f'<section aria-label="Book text"><h2>Chapter 1</h2><p>{PROSE}</p></section>'
UNUSUAL = f'<div class="new-publisher-wrapper"><h2>Chapter 1</h2><p>{PROSE}</p></div>'
EMBEDS = [
    '<iframe src="/reader/123" title="Silent Harbor"></iframe>',
    '<object data="/files/123.pdf" type="application/pdf"></object>',
    '<embed src="/files/123.pdf" type="application/pdf">',
    '<iframe src="https://external.invalid/reader/123" title="Silent Harbor"></iframe>',
    '<iframe srcdoc="' + escape(READER, quote=True) + '"></iframe>',
]


@pytest.fixture(autouse=True)
def no_live_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden in RT-05/06 regressions")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def card(content="", book=BOOK, heading="h1"):
    return (f'<article><{heading}>{book.title}</{heading}><p class="book-author">{book.author}</p>'
            '<p>Publication details: language English; edition 2026; publisher Harbor Press.</p>'
            f'{content}</article>')


def document(html, url=URL):
    return FetchResult(url, url, 200, {"Content-Type": "text/html; charset=utf-8"}, html.encode())


@pytest.mark.parametrize("name,sha256", ORIGINALS)
def test_original_audit_inputs(name, sha256):
    body = (FIXTURES / (name + ".html")).read_bytes()
    assert hashlib.sha256(body).hexdigest() == sha256
    result = evaluate_page(document(body.decode()), [BOOK], BOOK.id)
    assert result["status"] in {1, 2}  # Independent oracle does not require a positive.
    assert result["reason_code"] == "content_scope_unresolved"
    assert not result["content_signals"]
    assert result["match"]["unresolved_signals"]


@pytest.mark.parametrize("content", [UNUSUAL, SEMANTIC, SHORT, *EMBEDS])
def test_scoped_unverified_reading_stays_uncertain(content):
    result = evaluate_page(document(card(content)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (2, "content_scope_unresolved")
    assert not result["content_signals"]
    assert result["evidence"] and result["match"]["unresolved_signals"]


@pytest.mark.parametrize("content", [UNUSUAL, SEMANTIC, SHORT, *EMBEDS])
@pytest.mark.parametrize("outside_main", [False, True])
def test_independent_neighbour_cannot_change_target(content, outside_main):
    own, other = card(), card(content, OTHER, "h2")
    html = f'<main>{own}</main>{other}' if outside_main else f'<main>{own}{other}</main>'
    result = evaluate_page(document(html), [BOOK, OTHER], BOOK.id)
    assert (result["status"], result["reason_code"]) == (0, "content_not_confirmed")
    assert not result["content_signals"] and not result["match"]["unresolved_signals"]


@pytest.mark.parametrize("wrapper", [
    '<section class="annotation">{}</section>', '<section itemprop="description">{}</section>',
    '<section><h2>Book review</h2>{}</section>', '<section class="advertisement">{}</section>',
    '<div hidden>{}</div>', '<div style="display:none">{}</div>',
])
@pytest.mark.parametrize("content", [UNUSUAL, SEMANTIC, SHORT, *EMBEDS])
def test_commentary_advertising_and_hidden_reading_are_not_target_evidence(wrapper, content):
    result = evaluate_page(document(card(wrapper.format(content))), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (0, "content_not_confirmed")
    assert not result["content_signals"] and not result["match"]["unresolved_signals"]


@pytest.mark.parametrize("content", [
    "", f'<div class="description">{PROSE}</div>', f'<div class="review">{PROSE}</div>',
    f'<p>{PROSE}</p>', '<iframe src="/weather" title="Weather"></iframe>',
    '<iframe src="https://video.invalid/123"></iframe>', '<iframe></iframe>',
    '<object data="/logo.svg"></object>', '<embed src="/clip.mp4">',
    '<object data="javascript:book.pdf" type="application/pdf"></object>',
    '<iframe src="/viewer#book.pdf"></iframe>',
    f'<h2>{OTHER.title}</h2><p class="book-author">{OTHER.author}</p><h3>Chapter 1</h3><p>{PROSE}</p>',
])
def test_bibliography_prose_and_unrelated_embeds_remain_negative(content):
    result = evaluate_page(document(card(content)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (0, "content_not_confirmed")
    assert not result["content_signals"]


@pytest.mark.parametrize("content", [
    '<iframe src="/reader/123" title="Winter Garden"></iframe>',
    '<object data="/files/123.pdf" title="Winter Garden"></object>',
    '<embed src="/files/123.pdf" title="Winter Garden">',
])
def test_explicit_foreign_embedded_title_is_not_target_evidence(content):
    result = evaluate_page(document(card(content)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (0, "content_not_confirmed")
    assert not result["content_signals"] and not result["match"]["unresolved_signals"]


@pytest.mark.parametrize("proof", [READER, '<a href="/files/123.epub">Download EPUB</a>'])
@pytest.mark.parametrize("uncertain", [UNUSUAL, SHORT, *EMBEDS])
def test_sufficient_proof_keeps_one(proof, uncertain):
    result = evaluate_page(document(card(uncertain + proof)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")
    assert all(signal["binding"] == "bound" for signal in result["content_signals"])


def test_unverified_embeds_do_not_exhaust_positive_signal_budget():
    result = evaluate_page(document(card(EMBEDS[1] * 35 + READER)), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert len(result["extracted"]["observed_content_signals"]) <= 30


@pytest.fixture(scope="module")
def chromium():
    with BrowserProbe(timeout_ms=7000, fixture_routes={}) as probe:
        assert probe.start(), probe.failure
        yield probe


@pytest.mark.parametrize("content,status", [(READER, 1), (UNUSUAL, 2), (SEMANTIC, 2), (SHORT, 2),
                                           *[(embed, 2) for embed in EMBEDS], ("", 0)])
def test_real_chromium_keeps_scoped_decision(chromium, content, status):
    chromium.fixture_routes = {URL: (200, card(content))}
    if "srcdoc" in content:
        with chromium.browser.new_context(service_workers="block") as context:
            context.route("**/*", lambda route: route.abort())
            page = context.new_page()
            page.set_content(card(content))
            assert PROSE.strip() in page.frame_locator("iframe").locator(".reader-text").inner_text()
    rendered = chromium.check(URL)
    assert rendered.error_code == ""
    if "srcdoc" in content:
        # The main-document snapshot remains incomplete even when the frame rendered prose.
        assert "srcdoc=" in rendered.body.decode()
    result = evaluate_page(rendered, [BOOK], BOOK.id)
    assert result["status"] == status
    assert bool(result["content_signals"]) == (status == 1)


def test_public_cli_jsonl_xlsx_keep_original_decisions(tmp_path, monkeypatch):
    scenarios = [(name, (FIXTURES / (name + ".html")).read_text(encoding="utf-8"), 2)
                 for name, _ in ORIGINALS]
    scenarios += [("known", card(READER), 1), ("negative", card(), 0),
                  ("neighbour", card() + card(EMBEDS[1], OTHER, "h2"), 0)]
    source = tmp_path / "rt05_rt06.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet("Targets")
    sheet.append(HEADERS)
    responses = {}
    for name, html, _ in scenarios:
        url = URL + "/" + name
        sheet.append(("Publisher", BOOK.author, BOOK.title, "ebook", url, "offline.reader"))
        responses[url] = html.encode()
    workbook.save(source)
    workbook.close()
    original_hash = hashlib.sha256(source.read_bytes()).digest()
    _, targets = load_inputs(source)
    calls = []

    def synthetic_get(session, url, **kwargs):
        calls.append(url)
        response = requests.Response()
        response.status_code = 200
        response.url = url
        response.headers["Content-Type"] = "text/html; charset=utf-8"
        response.raw = BytesIO(responses[url])
        return response

    monkeypatch.setattr(requests.sessions.Session, "get", synthetic_get)
    output = tmp_path / "export"
    monkeypatch.setattr(sys, "argv", ["part2.py", "--input", str(source), "--browser", "off",
                                     "--delay", "0", "--output-dir", str(output)])
    main()
    records = [json.loads(line) for line in (output / "part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    workbook = openpyxl.load_workbook(output / "part2_results.xlsx", read_only=True)
    headers, *rows = list(workbook.active.values)
    workbook.close()
    assert calls == list(responses)  # No embedded resource was fetched.
    assert len(records) == len(rows) == len(scenarios)
    assert [record["status"] for record in records] == [case[2] for case in scenarios]
    for record, row, target in zip(records, rows, targets):
        cells = dict(zip(headers, row))
        for key in ["status", "reason_code", "expected_book_id", "source_row", "url", "method"]:
            assert record[key] == cells[key]
        assert record["expected_book_id"] == target["expected_book_id"]
        assert record["content_signals"] == json.loads(cells["content_signals"])
        assert record["match"] == json.loads(cells["catalog_match_details"])
    assert hashlib.sha256(source.read_bytes()).digest() == original_hash
