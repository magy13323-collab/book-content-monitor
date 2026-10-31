"""Strict е/ё candidates and uncertain spelling variants; synthetic offline data."""

import hashlib
import json

import openpyxl
import pytest
import requests

import part2
from pages import FetchResult, extract_page
from part1 import HEADERS, normalize_text, process_workbook, write_workbook
from part2 import Book, analyze_url, evaluate_page, load_inputs, load_json_catalog, write_results


AUTHOR = "Анна Лесная"
PAIRS = [("Все дома", "Всё дома"), ("Федор", "Фёдор")]
LEGACY_ID = "book_3a2c945a06f0ec9a"


@pytest.fixture(autouse=True)
def no_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def source(tmp_path, rows, name="input.xlsx"):
    path = tmp_path / name
    workbook = openpyxl.Workbook()
    for sheet in [workbook.active, workbook.create_sheet("Targets")]:
        sheet.append(HEADERS)
        for index, (author, title) in enumerate(rows):
            sheet.append(("Publisher", author, title, "ebook", f"https://offline.test/{index}", "offline.test"))
    workbook.save(path)
    workbook.close()
    return path


def document(title, author=AUTHOR, structured=None, structured_author=None, content=None):
    schema = '' if structured is None else '<script type="application/ld+json">' + json.dumps({
        "@type": "Book", "name": structured, "author": {"name": structured_author or author}}) + '</script>'
    content = '<a href="/files/123.epub">Скачать EPUB</a>' if content is None else content
    html = (f'<html><head><meta charset="utf-8">{schema}</head><body><main><h1>{title}</h1>'
            f'<p class="book-author">{author}</p><p>Карточка произведения с доступными форматами чтения.</p>'
            f'{content}</main></body></html>')
    return FetchResult("https://offline.test/book", "https://offline.test/book", 200,
                       {"Content-Type": "text/html; charset=utf-8"}, html.encode())


@pytest.mark.parametrize("rows", [
    [(AUTHOR, "Все дома"), (AUTHOR, "Всё дома")],
    [(AUTHOR, "Федор"), (AUTHOR, "Фёдор")],
    [("Федор Северный", "Дом"), ("Фёдор Северный", "Дом")],
])
def test_distinct_candidates_survive_both_excel_loaders_and_enter_review(tmp_path, rows):
    path = source(tmp_path, rows)
    analysis = process_workbook(path)
    catalog, targets = load_inputs(path)
    assert len(analysis.records) == len(analysis.books) == len(catalog) == 2
    assert len({book.id for book in catalog}) == 2
    assert [record.book_id for record in analysis.records] == [target["expected_book_id"] for target in targets]
    assert [(book.author, book.title) for book in catalog] == rows
    assert [record.raw[1:3] for record in analysis.records] == rows
    assert [record.source_row for record in analysis.records] == [2, 3]
    reviews = [item for item in analysis.review if item["code"] == "yo_variant"]
    assert len(reviews) == 1 and reviews[0]["source_rows"] == [2, 3]
    assert set(reviews[0]["book_ids"]) == {book.id for book in catalog}


def test_ids_are_stable_under_row_order_and_legacy_non_yo_id_survives(tmp_path):
    rows = [(AUTHOR, "Все дома"), (AUTHOR, "Всё дома")]
    first, _ = load_inputs(source(tmp_path, rows, "first.xlsx"))
    second, _ = load_inputs(source(tmp_path, list(reversed(rows)), "second.xlsx"))
    first_ids = {book.title: book.id for book in first}
    assert first_ids == {book.title: book.id for book in second}
    assert first_ids["Все дома"] == LEGACY_ID
    assert first_ids["Всё дома"] != LEGACY_ID


@pytest.mark.parametrize("title", ["Все дома", "Всё дома", "Фёдор"])
def test_exact_duplicates_keep_one_candidate(tmp_path, title):
    path = source(tmp_path, [(AUTHOR, title), (AUTHOR.upper(), '  ' + title.upper() + '  ')])
    analysis = process_workbook(path)
    catalog, targets = load_inputs(path)
    assert len(analysis.records) == 2 and len(analysis.books) == len(catalog) == 1
    assert targets[0]["expected_book_id"] == targets[1]["expected_book_id"]
    assert not any(item["code"] == "yo_variant" for item in analysis.review)


def test_loader_rejects_a_conflicting_id_instead_of_overwriting(tmp_path, monkeypatch):
    path = source(tmp_path, [(AUTHOR, "Дом"), (AUTHOR, "Лес")])
    monkeypatch.setattr(part2, "group_id", lambda *args: "collision")
    with pytest.raises(ValueError, match="Коллизия"):
        load_inputs(path)


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize("title", ["Все дома", "Всё дома", "Федор", "Фёдор"])
def test_exact_title_and_author_still_confirm_content(title, structured):
    book = Book("target", AUTHOR, title)
    result = evaluate_page(document(title, structured=title if structured else None), [book], book.id)
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize(("expected", "observed"), PAIRS + [(b, a) for a, b in PAIRS])
def test_title_spelling_alone_only_establishes_ambiguity(expected, observed, structured):
    book = Book("target", AUTHOR, expected)
    result = evaluate_page(document(observed, structured=observed if structured else None), [book], book.id)
    assert (result["status"], result["reason_code"]) == (2, "ambiguous")
    assert not result["content_signals"]
    assert any("е/ё" in item for item in result["evidence"])


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize(("expected", "observed"), [
    ("Федор Северный", "Фёдор Северный"), ("Фёдор Северный", "Северный Федор"),
    ("Анна Лесная; Федор Северный", "Фёдор Северный; Анна Лесная"),
])
def test_author_spelling_does_not_prove_identity(expected, observed, structured):
    book = Book("target", expected, "Дом")
    result = evaluate_page(document("Дом", observed, "Дом" if structured else None), [book], book.id)
    assert (result["status"], result["reason_code"]) == (2, "ambiguous")
    assert not result["content_signals"]


@pytest.mark.parametrize("observed", ["Фёдор Северный;", "Фёдор Северный;; Анна Лесная"])
def test_yo_similarity_does_not_bypass_coauthor_validation(observed):
    book = Book("target", "Федор Северный; Анна Лесная", "Дом")
    result = evaluate_page(document("Дом", observed), [book], book.id)
    assert (result["status"], result["reason_code"]) == (2, "author_conflict")


@pytest.mark.parametrize("title", ["Все дома", "Всё дома"])
def test_exact_catalogue_candidate_wins_without_confirming_its_variant(title):
    catalog = [Book("e", AUTHOR, "Все дома"), Book("yo", AUTHOR, "Всё дома")]
    result = evaluate_page(document(title), catalog)
    assert result["status"] == 1
    assert result["match"]["book_id"] == ("e" if title == "Все дома" else "yo")
    assert not result["match"]["other_catalog_matches"]
    other = "yo" if title == "Все дома" else "e"
    assert evaluate_page(document(title), catalog, other)["status"] == 2


@pytest.mark.parametrize(("dom", "schema"), [("Все дома", "Всё дома"), ("Всё дома", "Все дома")])
def test_jsonld_title_variant_cannot_inherit_the_dom_file(dom, schema):
    initial = document(dom, structured=schema)
    result = evaluate_page(initial, [Book("target", AUTHOR, schema)], "target")
    assert (result["status"], result["reason_code"]) == (2, "metadata_title_conflict")
    assert not next(s for s in extract_page(initial).subjects if s["source"] == "jsonld_book")["signals"]


def test_jsonld_author_variant_is_uncertain():
    initial = document("Дом", "Федор Северный", "Дом", "Фёдор Северный")
    result = evaluate_page(initial, [Book("target", "Федор Северный", "Дом")], "target")
    assert (result["status"], result["reason_code"]) == (2, "metadata_author_conflict")


def test_safe_decoration_does_not_hide_a_title_spelling_variant():
    result = evaluate_page(document("Читать «Все дома» онлайн"), [Book("target", AUTHOR, "Всё дома")], "target")
    assert (result["status"], result["reason_code"]) == (2, "ambiguous")


def test_variant_file_label_is_unresolved():
    initial = document("Все дома", content='<a href="/files/123.epub">Скачать Всё дома EPUB</a>')
    result = evaluate_page(initial, [Book("target", AUTHOR, "Все дома")], "target")
    assert (result["status"], result["reason_code"]) == (2, "content_scope_unresolved")


@pytest.mark.parametrize(("title", "filename"), [("Все дома", "Всё-дома"), ("Всё дома", "Все-дома")])
def test_complete_filename_spelling_variant_is_unresolved(title, filename):
    initial = document(title, content=f'<a href="/files/{filename}.epub">Скачать EPUB</a>')
    result = evaluate_page(initial, [Book("target", AUTHOR, title)], "target")
    assert (result["status"], result["reason_code"]) == (2, "content_scope_unresolved")
    assert not result["content_signals"]
    assert any("е/ё" in item for item in result["evidence"])


def test_nested_variant_file_is_uncertain_without_an_exact_own_file():
    initial = document("Все дома", content=f'<article><h2>Всё дома</h2><p class="book-author">{AUTHOR}</p>'
                       '<a href="/files/123.epub">Скачать EPUB</a></article>')
    result = evaluate_page(initial, [Book("target", AUTHOR, "Все дома")], "target")
    assert (result["status"], result["reason_code"]) == (2, "ambiguous")
    assert not result["content_signals"]
    assert any(s["binding"] == "unresolved" for s in result["extracted"]["observed_content_signals"])


def test_json_catalogue_keeps_both_candidates(tmp_path):
    path = tmp_path / "catalogue.json"
    path.write_text(json.dumps([{"author": AUTHOR, "title": title} for title in PAIRS[0]]), encoding="utf-8")
    catalog = load_json_catalog(path)
    assert len(catalog) == len({book.id for book in catalog}) == 2
    assert [book.title for book in catalog] == ["Все дома", "Всё дома"]


@pytest.mark.parametrize("crossed", [False, True], ids=["exact", "uncertain"])
def test_end_to_end_exports_preserve_both_rows_and_identities(tmp_path, crossed):
    path = source(tmp_path, [(AUTHOR, "Все дома"), (AUTHOR, "Всё дома")])
    original_hash = hashlib.sha256(path.read_bytes()).digest()
    analysis = process_workbook(path)
    catalog, targets = load_inputs(path)
    write_workbook(analysis, tmp_path / "processed.xlsx")
    results = []
    for target in targets:
        title = target["expected_title"]
        observed = ("Всё дома" if title == "Все дома" else "Все дома") if crossed else title
        initial = document(observed)
        results.append(analyze_url(target, catalog, fetcher=lambda url, initial=initial: initial))
    write_results(results, tmp_path / "output")
    exported = [json.loads(line) for line in (tmp_path / "output/part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["source_row"] for r in exported] == [2, 3]
    assert [r["expected_title"] for r in exported] == ["Все дома", "Всё дома"]
    assert len({r["expected_book_id"] for r in exported}) == 2
    assert [r["status"] for r in exported] == ([2, 2] if crossed else [1, 1])
    for name, sheet in [("processed.xlsx", "Normalized"), ("output/part2_results.xlsx", "Results")]:
        workbook = openpyxl.load_workbook(tmp_path / name, read_only=True)
        assert workbook[sheet].max_row == 3
        if sheet == "Normalized":
            rows = list(workbook[sheet].values)[1:]
            assert [(row[2], row[6]) for row in rows] == [("Все дома", 2), ("Всё дома", 3)]
            assert workbook["Books"].max_row == 3
            assert {row[10] for row in list(workbook["Books"].values)[1:]} == {"needs_review"}
        workbook.close()
    assert hashlib.sha256(path.read_bytes()).digest() == original_hash


def test_legacy_id_does_not_silently_resolve_to_a_new_yo_candidate(tmp_path):
    catalog, _ = load_inputs(source(tmp_path, [(AUTHOR, "Всё дома")]))
    result = evaluate_page(document("Всё дома"), catalog, LEGACY_ID)
    assert (result["status"], result["reason_code"]) == (2, "catalog_missing")


def test_unrelated_title_is_still_negative_and_search_normalization_is_unchanged():
    result = evaluate_page(document("Лес"), [Book("target", AUTHOR, "Дом")], "target")
    assert (result["status"], result["reason_code"]) == (0, "no_match")
    assert normalize_text("  ВСЁ\u00a0ДОМА ") == "все дома"
