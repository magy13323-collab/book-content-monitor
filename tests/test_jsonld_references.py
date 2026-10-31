"""Independent RT-02/03 oracles; complete HTML/PDF-free synthetic documents."""

from io import BytesIO
import json
import sys

import openpyxl
import pytest
import requests

from browser import BrowserProbe
from pages import FetchResult, extract_page
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main


BOOK = Book("target", "Alice North", "Silent Harbor")
OTHER = Book("other", "Bob West", "Winter Garden")
URL = "https://offline.jsonld/book"
FILE = '<a href="/files/123.epub">Download EPUB</a>'
TYPES = ["Book", "https://schema.org/Book", "http://schema.org/Book",
         ["Thing", "Book"], ["Thing", "https://schema.org/Book"]]


@pytest.fixture(autouse=True)
def no_live_http(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden in RT-02/03 regressions")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def script(data):
    return '<script type="application/ld+json">' + json.dumps(data) + "</script>"


def metadata(author=BOOK.author, kind="Book", identifier="#book", title=BOOK.title):
    return {"@context": "https://schema.org", "@type": kind, "@id": identifier,
            "name": title, "author": author if isinstance(author, (dict, list)) else {"name": author}}


def card(book=BOOK, author=True, content=FILE, structured=""):
    byline = f'<p class="book-author">{book.author}</p>' if author else ""
    return (f'<article id="{book.id}"><h1>{book.title}</h1>{byline}'
            '<p>Public bibliographic details, publisher and available editions.</p>' + structured + content + "</article>")


def document(data=None, body=None, author=True, scripts=None):
    head = "" if data is None else script(data)
    if scripts is not None:
        head = "".join(script(item) for item in scripts)
    html = '<html><head><meta charset="utf-8">' + head + '</head><body><main>' + (card(author=author) if body is None else body) + "</main></body></html>"
    return FetchResult(URL, URL, 200, {"Content-Type": "text/html; charset=utf-8"}, html.encode())


def graph(author=BOOK.author, kind="Book", nodes=None, author_id="#person"):
    return {"@context": "https://schema.org", "@graph": [
        {"@type": "WebPage", "@id": URL, "mainEntity": {"@id": "#book"}},
        metadata({"@id": author_id}, kind),
        *([{ "@type": "Person", "@id": "#person", "name": author}] if nodes is None else nodes),
    ]}


@pytest.mark.parametrize("kind", TYPES, ids=["short", "https", "http", "array-short", "array-uri"])
@pytest.mark.parametrize("conflict", [False, True], ids=["same-author", "conflict"])
@pytest.mark.parametrize("in_graph", [False, True], ids=["direct", "graph"])
def test_book_type_forms_preserve_author_conflicts(kind, conflict, in_graph):
    author = OTHER.author if conflict else BOOK.author
    data = metadata(author, kind)
    if in_graph:
        data = {"@graph": [{"@type": "WebPage", "@id": URL, "mainEntity": {"@id": "#book"}}, data]}
    result = evaluate_page(document(data), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == ((2, "metadata_author_conflict") if conflict else (1, "placement_signs"))
    subjects = [s for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book"]
    assert len(subjects) == 1 and subjects[0]["authors"] == [author]


@pytest.mark.parametrize("kind", TYPES, ids=["short", "https", "http", "array-short", "array-uri"])
def test_equivalent_types_can_supply_missing_dom_author(kind):
    result = evaluate_page(document(metadata(kind=kind), author=False), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert result["content_signals"]


@pytest.mark.parametrize("kind", TYPES, ids=["short", "https", "http", "array-short", "array-uri"])
@pytest.mark.parametrize("conflict", [False, True], ids=["same-author", "conflict"])
def test_unique_person_reference_preserves_conflict(kind, conflict):
    author = OTHER.author if conflict else BOOK.author
    result = evaluate_page(document(graph(author, kind)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == ((2, "metadata_author_conflict") if conflict else (1, "placement_signs"))
    assert next(s for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book")["authors"] == [author]


@pytest.mark.parametrize("person_id", ["#person", URL + "#person"])
@pytest.mark.parametrize("author_id", ["#person", URL + "#person"])
def test_relative_and_absolute_ids_resolve_the_same_named_person(person_id, author_id):
    data = graph(nodes=[{"@type": "https://schema.org/Person", "@id": person_id, "name": BOOK.author}], author_id=author_id)
    result = evaluate_page(document(data, author=False), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert result["match"]["state"] == "matched"


@pytest.mark.parametrize("reverse", [False, True])
@pytest.mark.parametrize("conflict", [False, True])
def test_reference_can_use_unique_definition_in_another_script(reverse, conflict):
    author = OTHER.author if conflict else BOOK.author
    scripts = [metadata({"@id": "#person"}), {"@type": "Person", "@id": "#person", "name": author}]
    result = evaluate_page(document(scripts=scripts[::-1] if reverse else scripts), [BOOK], BOOK.id)
    assert result["status"] == (2 if conflict else 1)
    assert next(s for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book")["authors"] == [author]


UNRESOLVED = [
    [],
    [{"@type": "Person", "@id": "#foreign", "name": BOOK.author}],
    [{"@type": "Person", "@id": "#person"}],
    [{"@type": "Person", "@id": "#person", "name": BOOK.author}, {"@type": "Person", "@id": "#person", "name": OTHER.author}],
    [{"@type": "Book", "@id": "#person", "name": BOOK.author}],
    [{"@type": "Person", "@id": "#person", "sameAs": {"@id": "#person"}}],
    [{"@type": "Person", "@id": "#person", "sameAs": {"@id": "#next"}}, {"@type": "Person", "@id": "#next", "sameAs": {"@id": "#person"}}],
]


@pytest.mark.parametrize("nodes", UNRESOLVED, ids=["missing", "foreign-id", "unnamed", "duplicate-conflict", "book-not-person", "self-cycle", "two-node-cycle"])
@pytest.mark.parametrize("dom_author", [False, True], ids=["no-dom-byline", "matching-dom-byline"])
def test_unresolved_or_ambiguous_reference_never_confirms(nodes, dom_author):
    result = evaluate_page(document(graph(nodes=nodes), author=dom_author), [BOOK], BOOK.id)
    assert result["status"] == 2
    assert result["reason_code"] == "author_missing"
    assert all(s["authors"] == [] for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book")


@pytest.mark.parametrize("author_id", ["", None, 42])
def test_malformed_explicit_id_is_not_an_author_confirmation(author_id):
    result = evaluate_page(document(metadata({"@id": author_id})), [BOOK], BOOK.id)
    assert result["status"] == 2


def test_incomplete_reference_list_is_not_a_complete_author_set():
    data = metadata([{"name": BOOK.author}, {"@id": "#missing"}])
    result = evaluate_page(document(data), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (2, "author_missing")


def test_unrelated_person_does_not_override_inline_author():
    data = {"@graph": [metadata(), {"@type": "Person", "@id": "#foreign", "name": OTHER.author}]}
    result = evaluate_page(document(data), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert next(s for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book")["authors"] == [BOOK.author]


def test_inline_person_with_its_own_unique_id_is_supported():
    result = evaluate_page(document(metadata({"@type": "Person", "@id": "#person", "name": BOOK.author}), author=False), [BOOK], BOOK.id)
    assert result["status"] == 1


@pytest.mark.parametrize("dom_author", [False, True])
def test_repeated_identical_person_definition_is_still_unambiguous(dom_author):
    nodes = [{"@type": "Person", "@id": "#person", "name": BOOK.author},
             {"@type": "Person", "@id": URL + "#person", "name": BOOK.author}]
    result = evaluate_page(document(graph(nodes=nodes), author=dom_author), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert next(s for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book")["authors"] == [BOOK.author]


def test_inline_name_does_not_bypass_conflicting_definitions_of_its_id():
    data = {"@graph": [metadata({"@id": "#person", "name": BOOK.author}), {"@id": "#person", "@type": "Person", "name": OTHER.author}]}
    result = evaluate_page(document(data), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (2, "author_missing")


@pytest.mark.parametrize("dom_author", [False, True])
def test_consistent_inline_and_graph_person_definitions_keep_the_positive(dom_author):
    data = {"@graph": [metadata({"@id": "#person", "name": BOOK.author}),
                       {"@id": "#person", "@type": "Person", "name": BOOK.author, "url": URL + "/author"}]}
    result = evaluate_page(document(data, author=dom_author), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert next(s for s in result["extracted"]["subjects"] if s["source"] == "jsonld_book")["authors"] == [BOOK.author]


def test_a_book_node_cannot_become_an_author_by_also_claiming_person_type():
    nodes = [{"@id": "#person", "@type": ["Book", "Person"], "name": BOOK.author}]
    result = evaluate_page(document(graph(nodes=nodes)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (2, "author_missing")


def test_consistent_person_name_variants_use_existing_author_equality():
    nodes = [{"@id": "#person", "@type": "Person", "name": BOOK.author},
             {"@id": "#person", "@type": "Person", "name": "North Alice"}]
    result = evaluate_page(document(graph(nodes=nodes), author=False), [BOOK], BOOK.id)
    assert result["status"] == 1


@pytest.mark.parametrize("name", [None, 42, {}, [BOOK.author]])
def test_non_string_person_name_cannot_become_a_confirmed_author(name):
    nodes = [{"@id": "#person", "@type": "Person", "name": name}]
    result = evaluate_page(document(graph(nodes=nodes)), [BOOK], BOOK.id)
    assert (result["status"], result["reason_code"]) == (2, "author_missing")


@pytest.mark.parametrize("kind", ["https://other.invalid/Book", ["Thing", "https://other.invalid/Book"]])
def test_foreign_type_namespace_is_not_treated_as_schema_book(kind):
    result = evaluate_page(document(metadata(OTHER.author, kind)), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert not any(s["source"] == "jsonld_book" for s in result["extracted"]["subjects"])


@pytest.mark.parametrize("broken_other", [False, True])
@pytest.mark.parametrize("reverse", [False, True])
def test_independent_card_references_and_issues_remain_local(broken_other, reverse):
    own_data = {"@graph": [metadata({"@id": "#own-person"}, "https://schema.org/Book"),
                           {"@type": "Person", "@id": "#own-person", "name": BOOK.author}]}
    other_data = {"@graph": [metadata({"@id": "#other-person"}, ["Book", "Thing"], identifier="#other-book", title=OTHER.title)]}
    if not broken_other:
        other_data["@graph"].append({"@type": "Person", "@id": "#other-person", "name": OTHER.author})
    cards = [card(structured=script(own_data)), card(OTHER, structured=script(other_data))]
    initial = document(body="".join(cards[::-1] if reverse else cards))
    own = evaluate_page(initial, [BOOK, OTHER], BOOK.id)
    other = evaluate_page(initial, [BOOK, OTHER], OTHER.id)
    assert own["status"] == 1
    assert all(s["scope_title"] == BOOK.title for s in own["content_signals"])
    assert other["status"] == (2 if broken_other else 1)


@pytest.mark.parametrize("linked", [False, True], ids=["uri-author", "linked-author"])
@pytest.mark.parametrize("target_conflict", [False, True], ids=["foreign-error", "target-error"])
def test_bm_open_16_conflicts_remain_in_their_independent_card(linked, target_conflict):
    cards = []
    for book, conflict in [(BOOK, target_conflict), (OTHER, not target_conflict)]:
        author = "Carol East" if conflict else book.author
        data = metadata({"@id": "#person-" + book.id} if linked else author,
                        "https://schema.org/Book", identifier="#" + book.id, title=book.title)
        if linked:
            data = {"@graph": [data, {"@id": "#person-" + book.id, "@type": "Person", "name": author}]}
        file = '<a href="/files/' + ("123" if book.id == BOOK.id else "456") + '.epub">Download EPUB</a>'
        cards.append(card(book, content=file, structured=script(data)))
    initial = document(body="".join(cards))
    own = evaluate_page(initial, [BOOK, OTHER], BOOK.id)
    other = evaluate_page(initial, [BOOK, OTHER], OTHER.id)
    assert (own["status"], other["status"]) == ((2, 1) if target_conflict else (1, 2))
    conflicting = own if target_conflict else other
    assert conflicting["reason_code"] == "metadata_author_conflict"
    assert [signal["url"] for signal in own["content_signals"]] == ["https://offline.jsonld/files/123.epub"]


def test_bm_open_16_unowned_reference_uncertainty_cannot_be_dismissed():
    other = Book("same-title", OTHER.author, BOOK.title)
    initial = document(metadata({"@id": "#missing"}), body=card() + card(other))
    for book in [BOOK, other]:
        result = evaluate_page(initial, [BOOK, other], book.id)
        assert (result["status"], result["reason_code"]) == (2, "author_missing")


@pytest.mark.parametrize("relation", ["mainEntity", "mainEntityOfPage"])
def test_main_entity_does_not_promote_unrelated_graph_book(relation):
    own = metadata({"@id": "#own-person"}, "https://schema.org/Book")
    page = {"@type": "WebPage", "@id": URL}
    if relation == "mainEntity":
        page[relation] = {"@id": "#book"}
    else:
        own[relation] = {"@id": URL}
    data = {"@graph": [page, own, {"@type": "Person", "@id": "#own-person", "name": BOOK.author},
                        metadata({"@id": "#foreign-person"}, identifier="#other-book", title=OTHER.title),
                        {"@type": "Person", "@id": "#foreign-person", "name": OTHER.author}]}
    result = evaluate_page(document(data), [BOOK], BOOK.id)
    assert result["status"] == 1
    assert "metadata_title_conflict" not in result["extracted"]["issues"]


@pytest.mark.parametrize("linked", [False, True])
def test_metadata_cannot_borrow_a_different_titles_file(linked):
    # Existing ownership invariant, without assuming an unavailable BM-OPEN label.
    longer = Book("longer", BOOK.author, BOOK.title + " and Beyond")
    data = graph(kind="https://schema.org/Book") if linked else metadata(kind="https://schema.org/Book")
    initial = document(data, body=card(longer))
    result = evaluate_page(initial, [BOOK], BOOK.id)
    assert result["status"] == 2
    assert not result["content_signals"]
    assert not next(s for s in extract_page(initial).subjects if s["source"] == "jsonld_book")["signals"]


@pytest.fixture(scope="module")
def chromium():
    pytest.importorskip("playwright.sync_api", reason="Optional Playwright unavailable")
    with BrowserProbe(timeout_ms=7000, fixture_routes={}) as probe:
        if not probe.start():
            pytest.skip(probe.failure)
        yield probe


@pytest.mark.parametrize(("data", "status"), [
    (metadata(OTHER.author, "https://schema.org/Book"), 2),
    (metadata(OTHER.author, ["Thing", "http://schema.org/Book"]), 2),
    (graph(OTHER.author), 2),
    (graph(), 1),
    (graph(nodes=[]), 2),
    (graph(nodes=UNRESOLVED[3]), 2),
], ids=["uri-conflict", "array-conflict", "linked-conflict", "linked-positive", "unresolved", "duplicate"])
def test_real_chromium_uses_the_same_safe_jsonld_rules(chromium, data, status):
    initial = document(data)
    chromium.fixture_routes = {URL: (200, initial.body.decode("utf-8"))}
    rendered = chromium.check(URL)
    assert rendered.error_code == "" and rendered.visibility_checked
    assert evaluate_page(rendered, [BOOK], BOOK.id)["status"] == status


def test_public_cli_exports_conflicts_and_unknowns_without_changing_ids(tmp_path, monkeypatch):
    scenarios = [(metadata(OTHER.author, "https://schema.org/Book"), 2),
                 (metadata(OTHER.author, "http://schema.org/Book"), 2),
                 (metadata(OTHER.author, ["Thing", "Book"]), 2),
                 (graph(OTHER.author), 2), (graph(), 1), (graph(nodes=[]), 2),
                 (graph(nodes=UNRESOLVED[3]), 2), (metadata(), 1)]
    source = tmp_path / "jsonld_input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet("Targets")
    sheet.append(HEADERS)
    responses = {}
    for index, (data, _) in enumerate(scenarios):
        url = URL + f"/{index}"
        sheet.append(("Publisher", BOOK.author, BOOK.title, "ebook", url, "offline.jsonld"))
        responses[url] = document(data).body
    workbook.save(source)
    workbook.close()
    catalog, targets = load_inputs(source)

    def synthetic_get(session, url, **kwargs):
        response = requests.Response()
        response.status_code = 200
        response.url = url
        response.headers["Content-Type"] = "text/html; charset=utf-8"
        response.headers["Content-Length"] = str(len(responses[url]))
        response.raw = BytesIO(responses[url])
        return response

    monkeypatch.setattr(requests.sessions.Session, "get", synthetic_get)
    output = tmp_path / "export"
    monkeypatch.setattr(sys, "argv", ["part2.py", "--input", str(source), "--browser", "off", "--delay", "0", "--output-dir", str(output)])
    main()
    lines = (output / "part2_results.jsonl").read_text(encoding="utf-8")
    records = [json.loads(line) for line in lines.splitlines()]
    workbook = openpyxl.load_workbook(output / "part2_results.xlsx", read_only=True)
    try:
        headers, *rows = list(workbook.active.values)
        cells = [dict(zip(headers, row)) for row in rows]
    finally:
        workbook.close()
    assert [r["status"] for r in records] == [status for _, status in scenarios]
    assert len(records) == len(cells) == len(targets) and len(catalog) == 1
    assert "_author_unresolved" not in lines
    for result, cell, target in zip(records, cells, targets):
        assert result["expected_book_id"] == target["expected_book_id"] == catalog[0].id
        for key in ["status", "reason_code", "source_row", "url", "expected_book_id", "method", "checked_at"]:
            assert result[key] == cell[key]
        for key in ["content_signals", "evidence", "observed_formats"]:
            assert result[key] == json.loads(cell[key])
