"""Literal title pipes and corroborated HTML site suffixes; no live requests."""

import hashlib
import json

import openpyxl
import pytest
import requests

from pages import FetchResult, extract_page
from part1 import process_workbook, write_workbook
from part2 import Book, analyze_url, evaluate_page, load_inputs, write_results
from test_part2 import synthetic_pdf
from test_yo_identity import source


AUTHOR = "Анна Лесная"
SUN, LONG, WIND = "Солнце", "Солнце | Луна", "Учебный маршрут | Том 1"


@pytest.fixture(autouse=True)
def no_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def document(heading, title=None, structured=None, site=None, og_title=None, content=None):
    metadata = '<meta charset="utf-8">'
    if title is not None:
        metadata += f'<title>{title}</title>'
    if site is not None:
        metadata += f'<meta property="og:site_name" content="{site}">'
    if og_title is not None:
        metadata += f'<meta property="og:title" content="{og_title}">'
    if structured is not None:
        schema = {"@type": "Book", "name": structured, "author": {"name": AUTHOR}}
        metadata += '<script type="application/ld+json">' + json.dumps(schema) + '</script>'
    identity = '' if heading is None else f'<h1>{heading}</h1>'
    content = '<a href="/files/123.epub">Скачать EPUB</a>' if content is None else content
    html = (f'<html><head>{metadata}</head><body><main>{identity}<p class="book-author">{AUTHOR}</p>'
            f'<p>Карточка произведения и доступные форматы для чтения книги.</p>{content}</main></body></html>')
    return FetchResult("https://offline.test/book", "https://offline.test/book", 200,
                       {"Content-Type": "text/html; charset=utf-8"}, html.encode())


@pytest.mark.parametrize("structured", [False, True], ids=["html", "jsonld"])
@pytest.mark.parametrize(("expected", "observed", "status", "reason"), [
    (SUN, LONG, 2, "ambiguous"),
    (WIND, WIND, 1, "placement_signs"),
    ("Учебный маршрут. Том 1", "Учебный маршрут. Том 1 | Том 2", 2, "volume_conflict"),
    (LONG, LONG, 1, "placement_signs"),
    (SUN, "Луна", 0, "no_match"),
    (LONG, "Солнце Луна", 2, "ambiguous"),
    ("Солнце Луна", LONG, 2, "ambiguous"),
    ("Солнце|", SUN, 2, "ambiguous"),
    (SUN, "|Солнце", 2, "ambiguous"),
    ("|Солнце", "|Солнце", 1, "placement_signs"),
    ("Солнце|", "Солнце|", 1, "placement_signs"),
])
def test_literal_h1_matrix(expected, observed, status, reason, structured):
    initial = document(observed, title=observed, structured=observed if structured else None)
    result = evaluate_page(initial, [Book("target", AUTHOR, expected)], "target")
    assert (result["status"], result["reason_code"]) == (status, reason)
    assert result["extracted"]["h1"] == [observed]
    if reason == "ambiguous":
        assert not result["content_signals"]


@pytest.mark.parametrize(("expected", "title", "site", "og_title", "schema", "status", "reason"), [
    (SUN, LONG, None, None, None, 2, "ambiguous"),
    (SUN, SUN, None, LONG, None, 2, "ambiguous"),
    (WIND, WIND, None, None, None, 1, "placement_signs"),
    (SUN, "Солнце | Библиотека", "Библиотека", None, None, 1, "placement_signs"),
    (SUN, LONG, "Библиотека", None, None, 2, "ambiguous"),
    (SUN, "Солнце | Библиотека", None, None, SUN, 2, "metadata_title_conflict"),
    (SUN, "Солнце | Библиотека", "Библиотека", None, SUN, 1, "placement_signs"),
    (WIND, WIND + " | Библиотека", "Библиотека", None, WIND, 1, "placement_signs"),
    (SUN, "Солнце | Библиотека", "Библиотека", "Солнце | Библиотека", None, 2, "ambiguous"),
    (SUN, "Солнце | Библиотека", None, None, None, 2, "ambiguous"),
])
def test_fallback_title_requires_independent_site_context(expected, title, site, og_title, schema, status, reason):
    initial = document(None, title, schema, site, og_title)
    result = evaluate_page(initial, [Book("target", AUTHOR, expected)], "target")
    assert (result["status"], result["reason_code"]) == (status, reason)
    assert result["extracted"]["title"] == title
    if site is not None:
        assert result["extracted"]["metadata"]["og:site_name"] == site


@pytest.mark.parametrize("title", [SUN, WIND])
def test_confirmed_h1_keeps_seo_html_title_out_of_the_book_name(title):
    initial = document(title, title + " | Название сайта")
    result = evaluate_page(initial, [Book("target", AUTHOR, title)], "target")
    assert (result["status"], result["reason_code"]) == (1, "placement_signs")


@pytest.mark.parametrize(("h1", "schema", "site", "status", "reason"), [
    (LONG, SUN, None, 2, "metadata_title_conflict"),
    (SUN, LONG, None, 2, "metadata_title_conflict"),
    (WIND, WIND, None, 1, "placement_signs"),
    ("Солнце | Библиотека", SUN, "Библиотека", 2, "metadata_title_conflict"),
])
def test_h1_and_jsonld_are_literal_book_titles(h1, schema, site, status, reason):
    initial = document(h1, structured=schema, site=site)
    result = evaluate_page(initial, [Book("target", AUTHOR, schema)], "target")
    assert (result["status"], result["reason_code"]) == (status, reason)
    if status == 2:
        assert not next(s for s in extract_page(initial).subjects if s["source"] == "jsonld_book")["signals"]


@pytest.mark.parametrize("extension", ["epub", "pdf"])
@pytest.mark.parametrize("selected", ["sun", "long"])
def test_independent_cards_keep_their_own_files(extension, selected):
    catalog = [Book("sun", AUTHOR, SUN), Book("long", AUTHOR, LONG)]
    content = ''.join(f'<article><h2>{b.title}</h2><p class="book-author">{AUTHOR}</p>'
                      f'<a href="/files/{b.id}.{extension}">Скачать EPUB</a></article>' for b in catalog)
    initial = document("Каталог книг", content=content)
    result = evaluate_page(initial, catalog, selected)
    assert result["status"] == 1
    assert [s["url"] for s in result["content_signals"]] == [f'https://offline.test/files/{selected}.{extension}']


def test_foreign_card_is_not_positive_evidence_for_a_shorter_title():
    content = (f'<article><h2>{SUN}</h2><p class="book-author">{AUTHOR}</p></article>'
               f'<article><h2>{LONG}</h2><p class="book-author">{AUTHOR}</p>'
               '<a href="/files/long.epub">Скачать EPUB</a></article>')
    result = evaluate_page(document("Каталог книг", content=content), [Book("sun", AUTHOR, SUN)], "sun")
    assert result["status"] != 1
    assert not result["content_signals"]


@pytest.mark.parametrize(("title", "label", "status", "reason"), [
    (LONG, LONG, 1, "placement_signs"),
    ("Солнце Луна", LONG, 2, "content_scope_unresolved"),
    (SUN, LONG, 2, "content_scope_unresolved"),
])
def test_file_label_must_identify_the_literal_title(title, label, status, reason):
    initial = document(title, content=f'<a href="/files/123.pdf">Скачать {label} PDF</a>')
    result = evaluate_page(initial, [Book("target", AUTHOR, title)], "target")
    assert (result["status"], result["reason_code"]) == (status, reason)


@pytest.mark.parametrize(("expected", "opening", "metadata", "status", "reason"), [
    ("North Wind | Volume 1", "North Wind | Volume 1", "North Wind | Volume 1", 1, "placement_signs"),
    ("Sun", "Sun | Moon", "Sun | Moon", 2, "ambiguous"),
    ("Sun Moon", "Sun | Moon", "Sun Moon", 2, "pdf_identity_unresolved"),
])
def test_pdf_titles_are_not_html_site_suffixes(expected, opening, metadata, status, reason):
    initial = FetchResult("https://offline.test/book.pdf", "https://offline.test/book.pdf", 200,
                          {"Content-Type": "application/pdf"}, synthetic_pdf(
                              title=opening, metadata={"/Title": metadata, "/Author": "Alice Example"}))
    result = evaluate_page(initial, [Book("target", "Alice Example", expected)], "target")
    assert (result["status"], result["reason_code"]) == (status, reason)


@pytest.mark.parametrize("crossed", [False, True], ids=["exact", "uncertain"])
def test_export_keeps_literal_names_and_stable_catalogue_ids(tmp_path, crossed):
    path = source(tmp_path, [(AUTHOR, SUN), (AUTHOR, LONG), (AUTHOR, WIND)])
    original = hashlib.sha256(path.read_bytes()).digest()
    analysis = process_workbook(path)
    catalog, targets = load_inputs(path)
    assert len(analysis.books) == len(catalog) == 3
    write_workbook(analysis, tmp_path / "processed.xlsx")
    results = []
    for target in targets:
        observed = LONG if crossed and target["expected_title"] == SUN else target["expected_title"]
        initial = document(observed)
        results.append(analyze_url(target, catalog, fetcher=lambda url, initial=initial: initial))
    write_results(results, tmp_path / "output")
    exported = [json.loads(line) for line in (tmp_path / "output/part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [r["source_row"] for r in exported] == [2, 3, 4]
    assert [r["expected_title"] for r in exported] == [SUN, LONG, WIND]
    assert [r["expected_book_id"] for r in exported] == [r.book_id for r in analysis.records]
    assert [r["status"] for r in exported] == ([2, 1, 1] if crossed else [1, 1, 1])
    workbook = openpyxl.load_workbook(tmp_path / "processed.xlsx", read_only=True)
    assert workbook["Normalized"].max_row == workbook["Books"].max_row == 4
    workbook.close()
    assert hashlib.sha256(path.read_bytes()).digest() == original
