"""Expected rules: lossless rows, conservative grouping, advisory review only."""

import hashlib
from pathlib import Path

import openpyxl
import pytest

from part1 import normalize_text, normalize_url, process_workbook, write_workbook


HEADERS = ["Праообладатель", "Автор", "Название", "Тип контента", "url", "Сайт"]
SOURCE = Path(__file__).resolve().parents[1] / "examples" / "synthetic_input.xlsx"


def make_source(tmp_path, rows):
    path = tmp_path / "input.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Исходные данные"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.create_sheet("Другой лист").append(["Не обрабатывать"])
    workbook.save(path)
    workbook.close()
    return path


def row(author="Автор", title="Книга", url="https://example.org/book", kind="ebook"):
    return ("Demo Rights", author, title, kind, url, "example.org")


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("  ЁЛКА\u00a0  и\tЕЖИ  ", "елка и ежи"),
        (" ;120 УЧЕБНЫХ МАРШРУТОВ; ", "120 учебных маршрутов"),
        ('«Книга» — Том III. Часть 2', '"книга" - том iii. часть 2'),
        ("(Не)полёт для робота", "(не)полет для робота"),
        ("(DemoAlias) Примерова Лада;", "(demoalias) примерова лада"),
        ("*7000 steps in a demo", "7000 steps in a demo"),
        ("Картаветра", "картаветра"),
        (None, ""),
    ],
)
def test_text_normalization_is_conservative_and_idempotent(value, expected):
    assert normalize_text(value) == expected
    assert normalize_text(normalize_text(value)) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (" HTTPS://Example.ORG:443/Book/?t=940730&x=2#Chapter ",
         "https://example.org/Book/?t=940730&x=2#Chapter"),
        ("http://EXAMPLE.org:80/?section_id=98", "http://example.org/?section_id=98"),
        ("https://forum.example/index.php?threads/11672/",
         "https://forum.example/index.php?threads/11672/"),
        ("https://s4.example.org/book/", "https://s4.example.org/book/"),
        ("https://example.org/%D0%81?q=a%2Fb&x=1&x=2",
         "https://example.org/%D0%81?q=a%2Fb&x=1&x=2"),
    ],
)
def test_url_normalization_preserves_resource_identity(value, expected):
    assert normalize_url(value) == expected
    assert normalize_url(normalize_url(value)) == expected


def test_import_and_export_preserve_values_and_physical_rows(tmp_path):
    original = row("  Ёж; ", " ;Книга «А» ", "https://EXAMPLE.org/Book?t=1")
    source = make_source(tmp_path, [original, (None,) * 6, row(title="Другая книга")])
    before = hashlib.sha256(source.read_bytes()).digest()
    analysis = process_workbook(source)
    assert len(analysis.records) == 2
    assert [record.source_row for record in analysis.records] == [2, 4]
    assert analysis.records[0].raw == original
    assert analysis.records[0].source_sheet == "Исходные данные"
    output = tmp_path / "processed.xlsx"
    write_workbook(analysis, output)
    workbook = openpyxl.load_workbook(output)
    sheet = workbook["Normalized"]
    assert sheet["A1"].value == "Правообладатель"
    assert tuple(sheet.cell(2, column).value for column in range(1, 7)) == original
    assert sheet.cell(2, 7).value == 2
    assert sheet.cell(3, 7).value == 4
    assert workbook.sheetnames == ["Normalized", "Duplicates", "Books", "Review"]
    assert all(sheet.auto_filter.ref and sheet.freeze_panes == "A2" for sheet in workbook)
    workbook.close()
    assert hashlib.sha256(source.read_bytes()).digest() == before


def test_full_duplicates_and_url_conflicts_are_different(tmp_path):
    source = make_source(tmp_path, [row(), row(), row(author="Другой автор")])
    analysis = process_workbook(source)
    first, second, third = analysis.records
    assert first.record_group_id == second.record_group_id != third.record_group_id
    assert first.url_group_id == second.url_group_id == third.url_group_id
    assert first.book_id == second.book_id != third.book_id
    duplicates = {item["type"]: item for item in analysis.duplicates}
    assert duplicates["full_record"]["source_rows"] == [2, 3]
    assert duplicates["repeated_url"]["source_rows"] == [2, 3, 4]
    assert any(item["code"] == "url_record_conflict" for item in analysis.review)


def test_same_url_with_conflicting_rightsholder_is_reviewed(tmp_path):
    other = ("Другой правообладатель",) + row()[1:]
    analysis = process_workbook(make_source(tmp_path, [row(), other]))
    assert len(analysis.books) == 1
    assert not any(item["type"] == "full_record" for item in analysis.duplicates)
    assert any(item["code"] == "url_record_conflict" for item in analysis.review)


def test_malformed_urls_are_preserved_and_reviewed(tmp_path):
    url = "https://example.org/bo\tok?t=1"
    analysis = process_workbook(make_source(tmp_path, [row(url=url)]))
    assert analysis.records[0].url == url
    assert "invalid_url" in analysis.records[0].issues


def test_multiple_urls_for_one_book_are_not_deleted(tmp_path):
    source = make_source(tmp_path, [row(), row(url="https://example.org/other")])
    analysis = process_workbook(source)
    assert len(analysis.records) == 2
    assert len(analysis.books) == 1
    assert analysis.books[0]["url_count"] == 2
    assert analysis.records[0].url_group_id != analysis.records[1].url_group_id
    assert [item["type"] for item in analysis.duplicates] == ["book_multiple_urls"]


@pytest.mark.parametrize(
    ("first", "second", "review_code"),
    [
        (row(author="Автор А"), row(author="Автор Б"), "title_author_conflict"),
        (row(title="Книга. Том 1"), row(title="Книга. Том 2"), None),
        (row(title="Книга. Часть 1"), row(title="Книга. Часть 2"), None),
        (row(title="Книга. Том III"), row(title="Книга. Том IV"), None),
        (row(author="Мирон Пример"), row(author="Пример Мирон"), "author_variant"),
        (row(author="Лев Макетный"), row(author="Лев Макетный (Лёв)"),
         "author_variant"),
        (row(title="Карта ветра"), row(title="Картаветра"), "title_spacing_variant"),
        (row(kind="ebook"), row(kind="audio"), "content_type_conflict"),
    ],
)
def test_ambiguous_or_distinct_books_are_not_automatically_merged(
    tmp_path, first, second, review_code
):
    source = make_source(tmp_path, [first, second])
    analysis = process_workbook(source)
    assert len(analysis.books) == 2
    assert analysis.records[0].book_id != analysis.records[1].book_id
    if review_code:
        assert any(item["code"] == review_code for item in analysis.review)


def test_mirrors_remain_distinct_urls(tmp_path):
    source = make_source(tmp_path, [
        row(url="https://example.org/book/"),
        row(url="https://s4.example.org/book/"),
    ])
    analysis = process_workbook(source)
    assert len({record.url_group_id for record in analysis.records}) == 2
    assert any(item["code"] == "possible_mirror" for item in analysis.review)


def test_url_volume_conflict_is_advisory_not_a_title_correction(tmp_path):
    title = 'Маршрут «Робот»#6. Архив "К"'
    source = make_source(tmp_path, [
        row(title=title, url="https://example.org/marshrut-robot-7-arhiv-k-ch-2/"),
    ])
    analysis = process_workbook(source)
    assert analysis.records[0].title == normalize_text(title)
    assert "url_number_conflict" in analysis.records[0].issues


@pytest.mark.parametrize(
    ("title", "url"),
    [
        ("Lumiverden", "https://example.org/books_lumiverden_67131402"),
        ("Луно-рецепты", "https://example.org/luno-recepty"),
        ("Марский зоркий твой", "https://example.org/marskiy-zorkiy-tvoy/"),
    ],
)
def test_legitimate_words_and_common_transliteration_do_not_trigger_false_hints(tmp_path, title, url):
    analysis = process_workbook(make_source(tmp_path, [row(title=title, url=url)]))
    assert "unsegmented_title" not in analysis.records[0].issues
    assert "title_url_hint" not in analysis.records[0].issues


def test_possible_missing_spaces_are_reviewed_without_repair(tmp_path):
    analysis = process_workbook(make_source(tmp_path, [
        row(title="Картасинегогорода", url="https://example.org/karta-sinego-goroda-123/"),
    ]))
    assert analysis.records[0].title == "картасинегогорода"
    assert "unsegmented_title" in analysis.records[0].issues


def test_missing_fields_do_not_create_a_shared_book_or_url(tmp_path):
    source = make_source(tmp_path, [row(author=None, url=None), row(author=None, url=None)])
    analysis = process_workbook(source)
    assert len(analysis.books) == 2
    assert len({record.url_group_id for record in analysis.records}) == 2
    assert all("missing_required" in record.issues for record in analysis.records)


def test_group_ids_do_not_depend_on_row_order(tmp_path):
    first = row(title="Первая книга")
    second = row(title="Вторая книга", url="https://example.org/second")
    path = make_source(tmp_path, [first, second])
    before = process_workbook(path)
    make_source(tmp_path, [second, first, row(title="Новая книга")])
    after = process_workbook(path)
    assert {record.title: record.book_id for record in before.records} == {
        record.title: record.book_id for record in after.records if record.title != "новая книга"
    }


def test_output_cannot_overwrite_input(tmp_path):
    source = make_source(tmp_path, [row()])
    before = source.read_bytes()
    with pytest.raises(ValueError, match="исходный"):
        write_workbook(process_workbook(source), source)
    assert source.read_bytes() == before


def test_synthetic_workbook_counts_and_known_review_cases(tmp_path):
    before = hashlib.sha256(SOURCE.read_bytes()).digest()
    analysis = process_workbook(SOURCE)
    assert len(analysis.records) == 18
    assert len({record.url for record in analysis.records}) == 15
    full_duplicates = [item for item in analysis.duplicates if item["type"] == "full_record"]
    assert len(full_duplicates) == 1
    assert sum(item["row_count"] - 1 for item in full_duplicates) == 1
    assert len(analysis.books) == 16
    assert len([item for item in analysis.duplicates if item["type"] == "book_multiple_urls"]) == 1
    records = {record.source_row: record for record in analysis.records}
    assert records[2].book_id == records[3].book_id == records[4].book_id
    assert len({records[index].book_id for index in [2, 5, 6, 7, 8]}) == 5
    assert "url_number_conflict" in records[7].issues
    assert "possible_mirror" in records[2].issues
    assert "title_url_hint" in records[17].issues
    assert "short_title_url" in records[16].issues
    assert "unsegmented_title" in records[15].issues
    assert {"author_variant", "title_spacing_variant", "missing_required", "invalid_url", "listing_url",
            "site_host_mismatch", "title_author_conflict", "content_type_conflict"} <= {
        item["code"] for item in analysis.review}
    assert all(item["reason"] and item["needed_data"] for item in analysis.review)
    output = tmp_path / "part1_processed.xlsx"
    write_workbook(analysis, output)
    original = openpyxl.load_workbook(SOURCE)
    exported = openpyxl.load_workbook(output)
    assert exported["Normalized"].max_row == 19
    for record, exported_row in zip(analysis.records, exported["Normalized"].iter_rows(min_row=2)):
        assert tuple(cell.value for cell in exported_row[:6]) == tuple(
            original.worksheets[0].cell(record.source_row, column).value for column in range(1, 7))
    original.close()
    exported.close()
    assert hashlib.sha256(SOURCE.read_bytes()).digest() == before
