"""RT-07: a bounded card scan cannot establish absence beyond its budget."""

from dataclasses import replace
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys

import openpyxl
import pytest
import requests

from browser import BrowserProbe
import pages
from pages import FetchResult, Limits, extract_page
from part1 import HEADERS
from part2 import Book, evaluate_page, load_inputs, main


BOOK = Book("target", "Alice North", "Silent Harbor")
OTHER = Book("other", "Bob West", "Winter Garden")
URL = "https://offline.cards/library"
INTRO = '<h1>Editorial notes</h1><p>Welcome to the reading room. Browse independent book cards below.</p>'
FIXTURE = Path(__file__).parent / "fixtures/rt07/card_enumeration_truncation_complete_main.html"
SHA256 = "97c68378fbb531aa639b3d9bf621e01b4d82df77b3529a239c799a690d0e4aaa"


@pytest.fixture(autouse=True)
def no_live_requests(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Live HTTP forbidden in RT-07 regressions")
    monkeypatch.setattr(requests.sessions.Session, "request", forbidden)


def card(book=BOOK, own_file=True, extra="", tag="article", error=None):
    file_id = 123 if book.id == BOOK.id else 999
    link = f'<a href="/files/{file_id}.epub">Download EPUB</a>' if own_file else ""
    metadata = ""
    if error:
        metadata = '<script type="application/ld+json">' + json.dumps({
            "@type": "Book", "name": "Unrelated title" if error == "title" else book.title,
            "author": {"name": "Unknown Writer" if error == "author" else book.author},
        }) + "</script>"
    return (f'<{tag}><h2>{book.title}</h2><p class="book-author">{book.author}</p>'
            f'<p>Publication details and available formats.</p>{metadata}{link}{extra}</{tag}>')


def foreign_cards(count, own_file=True):
    return [card(Book(f"foreign-{i}", "Foreign Writer", f"Foreign Work {i}"), own_file=own_file) for i in range(count)]


def document(cards, inside_main=False, intro=INTRO, url=URL):
    listing = "".join(cards)
    body = f'<main>{intro}{listing}</main>' if inside_main else f'<main>{intro}</main>{listing}'
    html = f'<html><head><meta charset="utf-8"><title>Reading room</title></head><body>{body}</body></html>'
    return FetchResult(url, url, 200, {"Content-Type": "text/html; charset=utf-8"}, html.encode())


def evaluate(cards, **kwargs):
    return evaluate_page(document(cards, **kwargs), [BOOK, OTHER], BOOK.id)


def test_original_audit_input_and_unchanged_oracle():
    body = FIXTURE.read_bytes()
    assert hashlib.sha256(body).hexdigest() == SHA256
    doc = FetchResult("https://audit.invalid/book", "https://audit.invalid/book", 200,
                      {"Content-Type": "text/html; charset=utf-8"}, body)
    result = evaluate_page(doc, [BOOK], BOOK.id)
    assert result["status"] in {1, 2}  # Original audit oracle permits either, never 0.
    assert (result["status"], result["reason_code"]) == (2, "partial_document")
    assert result["extracted"]["partial"] is True
    assert not result["content_signals"]


@pytest.mark.parametrize("foreign,position,status,partial", [
    (0, 0, 1, False),
    (149, 149, 1, False),
    (150, 150, 2, True),
    (150, 0, 1, True),
    (150, 74, 1, True),
    (150, 149, 1, True),
])
@pytest.mark.parametrize("inside_main", [False, True])
def test_card_order_and_budget_boundaries(foreign, position, status, partial, inside_main):
    cards = foreign_cards(foreign)
    cards.insert(position, card())
    result = evaluate(cards, inside_main=inside_main)
    assert result["status"] == status
    assert result["extracted"]["partial"] is partial
    assert [s["url"] for s in result["content_signals"]] == ([URL.rsplit("/", 1)[0] + "/files/123.epub"] if status == 1 else [])


@pytest.mark.parametrize("count,status,partial", [(150, 0, False), (151, 2, True)])
@pytest.mark.parametrize("identical", [False, True])
def test_no_target_requires_complete_scan(count, status, partial, identical):
    cards = [card(OTHER)] * count if identical else foreign_cards(count)
    result = evaluate(cards)
    assert (result["status"], result["reason_code"]) == (status, "partial_document" if partial else "no_match")
    assert result["extracted"]["partial"] is partial
    assert not result["content_signals"]


def test_duplicate_heading_candidates_do_not_consume_card_budget(monkeypatch):
    original = pages.content_signals
    visited = []

    def observed(node, *args, **kwargs):
        if node.name == "article":
            visited.append(id(node))
        return original(node, *args, **kwargs)

    monkeypatch.setattr(pages, "content_signals", observed)
    page = extract_page(document(foreign_cards(149) + [card()]))
    assert len(visited) == len(set(visited)) == 150
    assert len([s for s in page.subjects if s["source"] == "book_card"]) == 150
    assert page.partial is False  # 300 raw candidates identify only 150 actual DOM nodes.


def test_card_extraction_still_stops_at_150_unique_nodes(monkeypatch):
    original = pages.content_signals
    visited = []

    def observed(node, *args, **kwargs):
        if node.name == "article":
            visited.append(id(node))
        return original(node, *args, **kwargs)

    monkeypatch.setattr(pages, "content_signals", observed)
    page = extract_page(document(foreign_cards(180) + [card()]))
    assert len(visited) == len(set(visited)) == 150
    assert page.partial is True
    assert not any(s["title"] == BOOK.title for s in page.subjects)


@pytest.mark.parametrize("count,status", [(149, 1), (150, 2)])
def test_heading_only_card_uses_remaining_unique_budget(count, status):
    result = evaluate(foreign_cards(count) + [card(tag="div")])
    assert result["status"] == status
    assert result["extracted"]["partial"] is (status == 2)


@pytest.mark.parametrize("own_file,total,status", [
    (True, 150, 1), (False, 150, 0), (True, 151, 1), (False, 151, 2),
])
def test_neighbour_epubs_never_supply_target_evidence(own_file, total, status):
    result = evaluate([card(own_file=own_file), *foreign_cards(total - 1)])
    assert result["status"] == status
    assert [s["url"] for s in result["content_signals"]] == (["https://offline.cards/files/123.epub"] if own_file else [])


@pytest.mark.parametrize("own_file", [False, True])
def test_nested_foreign_card_cannot_lend_its_epub(own_file):
    result = evaluate([card(own_file=own_file, extra=card(OTHER)), *foreign_cards(149)])
    assert result["status"] == (1 if own_file else 2)
    assert result["extracted"]["partial"] is True
    assert [s["url"] for s in result["content_signals"]] == (["https://offline.cards/files/123.epub"] if own_file else [])


@pytest.mark.parametrize("target_error", [False, True])
@pytest.mark.parametrize("error", ["title", "author"])
def test_independent_jsonld_conflicts_keep_their_owner(target_error, error):
    result = evaluate([card(error=error if target_error else None),
                       card(OTHER, error=None if target_error else error), *foreign_cards(149)])
    assert result["status"] == (2 if target_error else 1)
    assert result["extracted"]["partial"] is True
    assert result["reason_code"] == ("metadata_" + error + "_conflict" if target_error else "placement_signs")
    assert all(s["url"] == "https://offline.cards/files/123.epub" for s in result["content_signals"])


def test_large_complete_main_is_not_automatically_partial():
    result = evaluate([], intro=INTRO + "<p>Editorial information about the library.</p>" * 1500)
    assert (result["status"], result["reason_code"]) == (0, "no_match")
    assert result["extracted"]["partial"] is False


def test_existing_text_limit_is_not_cleared_by_complete_card_scan():
    page = extract_page(document([card()]), replace(Limits(), max_text_chars=20))
    assert page.partial is True


@pytest.fixture(scope="module")
def chromium():
    with BrowserProbe(timeout_ms=7000, fixture_routes={}) as probe:
        assert probe.start(), probe.failure
        yield probe


@pytest.mark.parametrize("count,target_first,status,partial", [
    (149, False, 1, False), (150, False, 2, True), (150, True, 1, True),
])
def test_real_chromium_preserves_card_limit_decisions(chromium, count, target_first, status, partial):
    cards = foreign_cards(count)
    cards.insert(0 if target_first else len(cards), card())
    doc = document(cards)
    chromium.fixture_routes = {URL: (200, doc.body.decode())}
    rendered = chromium.check(URL)
    assert rendered.error_code == ""
    result = evaluate_page(rendered, [BOOK], BOOK.id)
    assert result["status"] == status
    assert result["extracted"]["partial"] is partial


def test_cli_jsonl_xlsx_preserve_partial_and_card_ownership(tmp_path, monkeypatch):
    cases = [
        ([card()], 1, False),
        ([*foreign_cards(149, own_file=False), card()], 1, False),
        ([*foreign_cards(150, own_file=False), card()], 2, True),
        (foreign_cards(150, own_file=False), 0, False),
        (foreign_cards(151, own_file=False), 2, True),
        ([card(), *foreign_cards(150, own_file=False)], 1, True),
        ([card(own_file=False), *foreign_cards(149, own_file=False)], 0, False),
        ([card(own_file=False, extra=card(OTHER)), *foreign_cards(149, own_file=False)], 2, True),
        (FIXTURE.read_bytes(), 2, True),
    ]
    source = tmp_path / "rt07.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(HEADERS)
    sheet = workbook.create_sheet("Targets")
    sheet.append(HEADERS)
    responses = {}
    for i, (cards, _, _) in enumerate(cases):
        url = URL + f"/{i}"
        sheet.append(("Publisher", BOOK.author, BOOK.title, "ebook", url, "offline.cards"))
        responses[url] = cards if isinstance(cards, bytes) else document(cards, url=url).body
    workbook.save(source)
    workbook.close()
    source_sha = hashlib.sha256(source.read_bytes()).digest()
    _, targets = load_inputs(source)
    calls = []

    def synthetic_get(session, url, **kwargs):
        calls.append(url)
        response = requests.Response()
        response.status_code = 200
        response.url = url
        response.headers["Content-Type"] = "text/html; charset=utf-8"
        response.raw = BytesIO(responses[url])
        return response

    monkeypatch.setattr(requests.sessions.Session, "get", synthetic_get)
    output = tmp_path / "export"
    monkeypatch.setattr(sys, "argv", ["part2.py", "--input", str(source), "--browser", "off",
                                     "--delay", "0", "--output-dir", str(output)])
    main()
    records = [json.loads(line) for line in (output / "part2_results.jsonl").read_text(encoding="utf-8").splitlines()]
    workbook = openpyxl.load_workbook(output / "part2_results.xlsx", read_only=True)
    headers, *rows = list(workbook.active.values)
    workbook.close()
    assert calls == list(responses)
    assert len(records) == len(rows) == len(cases)
    for record, row, target, (_, status, partial) in zip(records, rows, targets, cases):
        cells = dict(zip(headers, row))
        assert record["status"] == status and record["extracted"]["partial"] is partial
        assert record["expected_book_id"] == target["expected_book_id"]
        for key in ["status", "reason_code", "expected_book_id", "url", "source_row", "method"]:
            assert record[key] == cells[key]
        assert cells["partial"] is partial
        assert record["content_signals"] == json.loads(cells["content_signals"])
        assert record["match"] == json.loads(cells["catalog_match_details"])
        assert all(s["url"] == "https://offline.cards/files/123.epub" for s in record["content_signals"])
    assert hashlib.sha256(source.read_bytes()).digest() == source_sha
